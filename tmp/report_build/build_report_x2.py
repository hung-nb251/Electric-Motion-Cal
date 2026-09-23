from pathlib import Path
import sys, os, math, json, re
sys.path.insert(0,str(Path(__file__).resolve().parent/'pydeps'))
os.environ['MPLCONFIGDIR']=str(Path(__file__).resolve().parent/'mplconfig')
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'outputs/bao_cao_xe_con'; TMP=ROOT/'tmp/report_build'
res=json.loads((TMP/'x2_results.json').read_text(encoding='utf-8'))
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10})
arr=np.array([row[:7] for row in res['graph']],dtype=float)
fig,axs=plt.subplots(4,1,figsize=(8,6.2),sharex=True)
for ax,idx,label in zip(axs,[1,2,3,4],['v (m/s)','F (N)','P motor (W)','M motor (N·m)']):
    ax.plot(arr[:,0],arr[:,idx],color='#1F4E78',lw=1.6)
    ax.axvspan(92,122,color='#dddddd',alpha=.6)
    ax.set_ylabel(label);ax.grid(alpha=.25);ax.axhline(0,color='#888888',lw=.6)
axs[0].set_title('Có tải đi (0–92 s)  |  Nghỉ (92–122 s)  |  Không tải về (122–211,5 s)',fontsize=10)
axs[-1].set_xlabel('Thời gian trong chu kỳ (s)');fig.tight_layout()
fig.savefig(TMP/'x2_profile.png',dpi=190);plt.close(fig)

# Reuse only the document formatting helpers and unchanged safety narrative.
old=(TMP/'build_report.py').read_text(encoding='utf-8')
exec(old[old.index('d=Document();'):old.index("d.add_paragraph('Thiết kế phần cứng")])
d.add_paragraph('Thiết kế hệ truyền động xe con cầu trục — x = 2',style='Title')
p('Báo cáo tính toán phần cứng và logic trạng thái theo bảng mẫu khóa trước',style='Subtitle')
h('1 Đầu vào và cơ sở tính toán')
p('Đối tượng thiết kế là truyền động di chuyển ngang của xe con cầu trục. Bản này thay các chữ số x bằng 2, sử dụng cấu hình cơ sở của sheet “Tính toán chi tiết động cơ a=1”, đồng thời sửa các công thức chưa nhất quán trong mẫu. Kết quả chọn sơ bộ động cơ 5,5 kW; không sử dụng phương án cản lăn 0,01 và motor 2,2 kW của bản trước.')
tab(['Đại lượng','Quy luật theo x','x = 2'],[['Tải định mức Q','10 000 + 1 111x','12 222 kg'],['Khối lượng cơ cấu nâng hạ m_h','200 + 11x','222 kg'],['Chiều cao nâng h','10 + x','12 m'],['Tốc độ ngang tối đa','0,1 + 0,01x','0,12 m/s'],['Gia tốc cho phép','Không đổi','≤ 0,5 m/s²'],['Độ giật cho phép','Không đổi','≤ 2 m/s³']],[2.7,2.3,2])
p('Chiều cao nâng 12 m không phải hành trình ngang. Khối lượng chuyển động có tải theo dữ liệu đề là 12 444 kg; không tải là 222 kg. Chưa có khối lượng khung xe và phụ kiện chưa nằm trong m_h; ô khối lượng bổ sung để 0 nhằm tái hiện dữ liệu đề, phải nhập số thực khi chốt thiết kế.')
tab(['Giả thiết kế thừa mẫu x = 1','Giá trị'],[['Đường kính bánh D; tỷ số truyền i','0,25 m; 100'],['Hệ số cản μ; hiệu suất cơ khí η','0,2; 0,96'],['Hệ số dự trữ k; gia tốc trọng trường g','1,2; 9,81 m/s²'],['Quán tính rotor motor tham khảo','0,039 kg·m²']],[4.4,2.6])
p('Các sheet của mẫu đúng là tương ứng x = 1, x = 3 và x = 6. Sheet a=6 có ô E5 ghi 1 nhưng các giá trị tải, khối lượng, chiều cao và tốc độ là của x=6. Tỷ số truyền, hiệu suất và mã motor là lựa chọn thiết kế của từng sheet, không phải đại lượng nội suy theo x.')
p('μ = 0,2 được giữ để tính đúng mô hình bài tập. Đây là hệ số cản trong công thức mẫu; không tự đồng nhất với hệ số bám bánh–ray. Kết quả công suất chưa chứng minh cơ cấu thực đủ bám hoặc đủ bền.')

page();h('2 Lực kéo và quy đổi về trục động cơ')
eq('F_đều = (Q + m_h) μ g = 24 415,128 N')
eq('F_tăng = (Q + m_h) (μ g + a)')
eq('F_giảm = (Q + m_h) (μ g − a₂)')
p('Trong công thức giảm tốc, a₂ là độ lớn gia tốc chậm dần. Nếu dùng gia tốc có dấu a < 0, viết F = m(μg + a). Khi chuyển sang lượt về, đổi dấu lực và tốc độ theo trục tọa độ chung.')
tab(['Đại lượng','Phép tính','Kết quả'],[['Lực ở giới hạn a = 0,5','12 444 × (0,2 × 9,81 + 0,5)','30 637,128 N'],['Mô-men tổng tại bánh','F × D/2','3 829,641 N·m'],['Mô-men quy đổi có dự trữ','k × F × D/(2iη)','47,8705 N·m'],['Tốc độ bánh','60v/(πD)','9,1673 vòng/phút'],['Tốc độ motor yêu cầu','i × n_bánh','916,7325 vòng/phút'],['Tốc độ góc motor','2vi/D','96 rad/s'],['Công suất cơ có dự trữ','M_motor × ω_motor','4 595,5692 W']],[2.65,2.45,1.9])
p('Hiệu suất đã nằm trong mô-men quy đổi, do đó công suất cơ trên trục motor là P = Mω. Nếu chia thêm η như ô E24 của mẫu, kết quả thành 4 787,0513 W. Bản x2 sửa lỗi tính hiệu suất hai lần; cả hai con số vẫn dẫn đến cấp motor 5,5 kW theo phép chọn công suất giản lược.')
sub('Quy đổi quán tính và bổ sung mô-men tăng tốc rotor')
eq('J_tải = (Q + m_h) [D/(2i)]² = 0,01944375 kg·m²')
eq('J_động lực = J_rot + J_phanh + J_bánh/i² + J_tải/η')
eq('M_motor = F D/(2iη) + (J_rot + J_phanh + J_bánh/i²) α_motor')
eq('α_motor = (2i/D) a;   P_motor = M_motor ω_motor')
p('J_tải là quán tính quy đổi động học; J_động lực có thêm hiệu suất để biểu diễn mô-men yêu cầu trong mô hình truyền thuận. Đã có ma trong F thì không cộng lại J_tải α/η lần nữa. Quán tính phanh và bánh chưa được cung cấp; các ô tương ứng để 0 và ghi rõ cần bổ sung.')

page();h('3 Chu kỳ phụ tải và giới hạn độ giật')
p('Giữ các thời gian chính của mẫu: tăng tốc 2 s; có tải chạy đều 80 s, giảm tốc 10 s; nghỉ 30 s; không tải tăng tốc 2 s, giảm tốc 5 s. Để xe về đúng điểm xuất phát, thời gian chạy đều không tải được sửa từ 80 s thành 82,5 s.')
tab(['Giai đoạn','Thời gian (s)','Quãng đường (m)'],[['Có tải: tăng tốc','2','0,12'],['Có tải: chạy đều','80','9,60'],['Có tải: giảm tốc','10','0,60'],['Nghỉ, motor bỏ mô-men','30','0'],['Không tải: tăng tốc','2','0,12'],['Không tải: chạy đều','82,5','9,90'],['Không tải: giảm tốc','5','0,30'],['Toàn chu kỳ','211,5','10,32 đi + 10,32 về']],[3.7,1.65,1.65])
p('Nếu cùng giữ 80 s chạy đều cho cả hai lượt, mẫu cho 10,32 m đi và 10,02 m về; hai lượt không khép kín. Hành trình 10,32 m là hệ quả của giả thiết thời gian trên, không phải h = 12 m. Chu kỳ chỉ có một khoảng nghỉ 30 s theo mẫu; khi bổ sung thời gian lấy tải ở đầu còn lại phải tính lại RMS.')
sub('Thay ramp tuyến tính bằng S-curve')
eq('t_j = 0,25 s;   a_p = v_max/(T_ramp − t_j);   j_p = a_p/t_j')
tab(['Ramp','|a| đỉnh (m/s²)','|j| đỉnh (m/s³)'],[['Tăng tốc 2 s','0,068571','0,274286'],['Giảm tốc có tải 10 s','0,012308','0,049231'],['Giảm tốc không tải 5 s','0,025263','0,101053']],[3.3,1.85,1.85])
p('Mỗi ramp gồm ba pha: jerk không đổi trong 0,25 s, gia tốc không đổi trong T_ramp − 0,5 s, rồi jerk ngược dấu trong 0,25 s. Gia tốc liên tục, đạt đúng vận tốc cuối; a và j đều thấp hơn giới hạn đề bài. Vận hành không cần đạt a = 0,5, vì đây là giới hạn trên.')
p('Tại các đoạn cột M của sheet Tinh_toan_x2, vẫn giữ phép tính ramp tuyến tính để đối chiếu mẫu. Kết quả chọn động cơ theo chu kỳ và đồ thị cuối cùng lấy từ sheet Chu_ky dùng S-curve; không trộn hai quỹ đạo.')

page();h('4 Đồ thị v(t), F(t), P(t), M(t) và tải hiệu dụng')
pic('x2_profile.png',6.7)
p('Hình 1. Chu kỳ có tải đi và không tải về. F là lực tại bánh, P và M là công suất và mô-men tại trục motor. Chiều về có v, F, M âm; công suất kéo vẫn dương khi lực cùng chiều vận tốc. Mô-men ma sát được mô hình hóa theo từng pha chuyển động.',style='Caption')
eq('P_hd = √[ ∫ P_motor²(t) dt / T_ck ] = 1 924,6695 W')
eq('M_RMS = √[ ∫ M_motor²(t) dt / T_ck ] = 20,9738 N·m')
eq('M_đỉnh = 35,0410 N·m;   k P_hd = 2 309,6034 W')
p('Bảng tính tích phân bình phương theo từng pha, kể cả nghỉ với P = M = 0. Trong mỗi pha jerk không đổi, P² là đa thức bậc tối đa 6; tích phân Gauss 4 điểm cho kết quả chính xác theo mô hình này. Không lấy bình phương công suất trung bình của pha để thay cho trung bình bình phương công suất.')
p('P_hd phục vụ tiêu chí công suất hiệu dụng trong đề; M_RMS hỗ trợ đánh giá tải nhiệt khi tốc độ thay đổi. Điều kiện nhiệt thực còn phụ thuộc làm mát, tổn hao không tải, tần suất thao tác và chế độ làm việc của motor.')

page();h('5 Chọn động cơ và kiểm nghiệm mô-men')
p('Chọn sơ bộ motor không đồng bộ ba pha 5,5 kW, 6 cực, tham khảo M3AA 132ME 6 đã ghi trong sheet x=1. Các trị số nhãn dưới đây được kế thừa bảng mẫu; chưa thay cho catalog xác nhận đúng mã đặt hàng.')
tab(['Thông số motor kế thừa','Giá trị'],[['Công suất, điện áp, tần số','5,5 kW; 400 V; 50 Hz'],['Tốc độ, mô-men danh định','973 vòng/phút; 53,9 N·m'],['Dòng định mức; cosφ','12 A; 0,74'],['Mô-men khởi động trực tiếp','2 × 53,9 = 107,8 N·m'],['Quán tính rotor','0,039 kg·m²']],[4.4,2.6])
tab(['Kiểm tra','Yêu cầu tính toán','Đánh giá sơ bộ'],[['Công suất theo mẫu','5,5 ≥ 4,5956 kW','Đạt'],['Hiệu dụng có dự trữ','5,5 ≥ 2,3096 kW','Đạt'],['Tốc độ làm việc','916,73 < 973 vòng/phút','VFD điều chỉnh tốc độ'],['Mô-men quỹ đạo có dự trữ','53,9 ≥ 1,2 × 35,041 = 42,049 N·m','Đạt'],['Mô-men hiệu dụng','20,974 < 53,9 N·m','Đạt so sánh số; cần kiểm tra nhiệt'],['Bao mô-men ở a = 0,5','66,591 > 53,9 N·m','Cần quá tải ngắn hạn ≥ 1,235 lần']],[2.1,2.75,2.15])
p('Bao mô-men ở a = 0,5 gồm mô-men tăng tốc rotor: k × [39,8921 + 0,039 × 800 × 0,5] = 66,5905 N·m. Không được kết luận motor chịu liên tục toàn bộ bao này chỉ từ công suất 4,60 kW. Quỹ đạo tăng tốc 2 s đã chọn chỉ cần 42,05 N·m sau dự trữ, nhỏ hơn 53,9 N·m.')
p('Mô-men khởi động trực tiếp 107,8 N·m lớn hơn bao yêu cầu 66,59 N·m về mặt số học; tuy nhiên chạy VFD phải dùng giới hạn dòng/mô-men của VFD và đường đặc tính motor. Cần kiểm tra khả năng quá tải ở khoảng 917 vòng/phút nếu tăng gia tốc đặt. Không dùng hệ số khởi động trực tiếp để bảo đảm mô-men VFD.')
p('Ở công suất 5,5 kW và tốc độ 973 vòng/phút, mô-men tính từ P/ω khoảng 53,98 N·m; 53,9 N·m của mẫu được giữ làm giá trị kiểm tra bảo thủ. Cần encoder và điều khiển phù hợp để giữ vận tốc ngang không vượt 0,12 m/s.')

page();h('6 Thông số cơ khí sơ bộ')
tab(['Hạng mục','Giá trị và điều kiện'],[['Bánh xe','D = 250 mm theo mẫu; số bánh và phân bố tải cần xác nhận.'],['Hộp giảm tốc','i = 100; η = 0,96 theo mẫu.'],['Mô-men tổng đầu ra do tải','M = F_max × 0,125 = 3 829,641 N·m.'],['Mô-men đầu ra có dự trữ','kM = 4 595,569 N·m = 4,596 kN·m.'],['Chia cho hai nhánh đều nhau','Mỗi nhánh ≥ 2 297,785 N·m; chỉ đúng khi thực sự chia tải đều.'],['Phản lực tĩnh nếu có bốn bánh đều tải','12 444 × 9,81 / 4 = 30,519 kN/bánh.'],['Tải bánh khảo sát với hệ số 1,25','38,149 kN/bánh; chọn cấp ≥ 40 kN để kiểm tra tiếp.'],['Trục chung theo xoắn thuần','Với [τ] giả định 40 MPa: d_min ≈ 83,6 mm; khảo sát d = 90 mm.']],[2.4,4.6])
eq('d_min = ∛[16 M/(π [τ])]  (M tính bằng N·mm)')
p('Đường kính trục trên là bước tính sơ bộ. Cần kiểm tra uốn–xoắn, then, tập trung ứng suất, độ võng và ổ lăn. Chưa có thông số ray, bề rộng bánh, vật liệu và độ cứng nên chưa kiểm tra tiếp xúc bánh–ray hoặc tuổi thọ ổ.')
sub('Giới hạn mô-men truyền động')
p('Motor có thể phát mô-men lớn hơn nhu cầu kéo. Hộp số, khớp nối và trục phải được kiểm tra thêm theo mô-men cực đại VFD, tình huống kẹt bánh và phanh; định mức đầu ra 4,596 kN·m chỉ đáp ứng tải tính toán có dự trữ, không mặc nhiên chịu mọi mức quá tải motor.')
sub('Điều kiện bám phải kiểm chứng riêng')
eq('F_kéo ≤ μ_bám N_chủ động')
p('Giữ μ cản = 0,2 theo mẫu không có nghĩa μ_bám = 0,2. Nếu giả sử hai bánh chủ động chịu 50% tổng tải và μ_bám = 0,2 thì lực bám chỉ 12 207,564 N, thấp hơn cả lực cản đều 24 415,128 N. Do đó không thể đồng thời dùng các giả thiết này để khẳng định cơ cấu thực chạy được. Cần xác nhận mô hình cản với giảng viên và dữ liệu bánh–ray trước khi chuyển từ bài tính sang thiết kế chế tạo.')
p('Khung xe, hộp số, motor và phụ kiện thực làm tăng tải bánh và lực kéo. Những bộ phận chưa nằm trong 222 kg phải cộng vào ô E51 của bảng tính, không tính trùng khối lượng.')

page();h('7 Biến tần, bảo vệ và hãm')
tab(['Thiết bị','Tính theo tiêu chí đề','Lựa chọn sơ bộ'],[['Biến tần','P_VFD ≥ 1,3 × 5,5 = 7,15 kW','Cấp 7,5 kW; xác minh dòng HD, quá tải, STO và chopper.'],['MCCB/cầu chì','I ≥ 1,5 × 12 = 18 A','Mốc 20 A theo bài tập; chốt lại theo đầu vào VFD và phối hợp bảo vệ.'],['Phanh cơ','Đóng bằng lò xo, nhả bằng điện','Cần phản hồi phanh và dữ liệu mô-men, thời gian đóng, khả năng hãm động.'],['Điện trở hãm','R ≥ R_min,VFD; kiểm tra năng lượng và công suất xung','Chưa chốt Ω/W khi chưa có mã VFD/chopper.']],[1.4,2.6,3])
p('Quy tắc 1,3–1,5 trong đề cho khoảng công suất biến tần 7,15–8,25 kW. Chọn hệ số 1,3 để đề xuất cấp 7,5 kW; nếu yêu cầu hệ số 1,5, phải chọn cấp ≥ 8,25 kW. Dòng liên tục và quá tải sau giảm định mức vẫn là điều kiện bắt buộc; không chỉ so sánh kW.')
eq('E_k,có tải = 0,5 m v² + 0,5 J_rot ω_motor² = 269,3088 J')
eq('E_thiết kế = 1,2 E_k = 323,1706 J')
eq('P_R,xung = U_chopper²/R;   E_R = ∫ P_R(t) dt')
p('Động năng là mức trên sơ bộ của năng lượng phải tiêu tán khi bỏ qua ma sát trên ray ngang, chưa gồm dao động tải và ngoại lực. Không lấy toàn bộ 323,17 J làm năng lượng tái sinh chắc chắn: phần lớn có thể tiêu tán trong cản chạy. Để chọn điện trở cần R_min, ngưỡng chopper, công suất xung cho phép và chu kỳ hãm từ catalog VFD.')
sub('Khoảng cách dừng theo quỹ đạo đã chọn')
eq('s_dừng,có tải = 0,5 × 0,12 × 10 = 0,60 m')
eq('d_kích hoạt ≥ v_max t_trễ + s_dừng + Δs')
p('Với giả thiết tổng trễ 0,10 s và dự trữ 0,05 m: d ≥ 0,662 m; lấy mốc khảo sát 0,70 m trước điểm dừng. Không tải giảm tốc 5 s cần 0,30 m. Đây là dừng vận hành từ tốc độ đều; Stop giữa ramp phải tính quỹ đạo từ trạng thái hiện tại. Cực hạn an toàn bố trí theo quãng đường dừng khẩn thực tế, không dùng 0,70 m như giá trị đã chứng nhận.')
p('Mạch động lực: nguồn → cách ly/MCCB → VFD → motor → hộp số → bánh xe. PLC/HMI xử lý quỹ đạo và trình tự; nhánh relay/PLC an toàn xử lý E-stop, cực hạn, STO và cắt điện nhả phanh độc lập. Phanh dùng dừng khẩn phải được phép hãm động theo dữ liệu hãng.')

safety=old[old.index("page();h('10 Tín hiệu"):old.index("page();h('14 Điều kiện")]
for a,b in [('10 Tín hiệu','8 Tín hiệu'),('11 Lưu đồ','9 Lưu đồ'),('12 Bảng chuyển','10 Bảng chuyển'),('13 Quy trình','11 Quy trình'),('Hình 3.','Hình 2.')]:safety=safety.replace(a,b)
exec(safety)

page();h('12 Mức hoàn thiện và tài liệu tham khảo')
p('Bản x2 đã có phép tính lực, quy đổi mô-men/quán tính, chu kỳ khép kín, đồ thị v–F–P–M, công suất hiệu dụng, kiểm nghiệm motor sơ bộ, tiêu chí biến tần/bảo vệ và lưu đồ trạng thái. Động cơ được đề xuất 5,5 kW cho quỹ đạo vận hành đã chọn, với các giả thiết của mẫu.')
p('Chưa chốt thiết bị chế tạo: cần xác nhận khối lượng và quán tính còn thiếu; lực cản và bám ray; catalog motor và chế độ nhiệt; hộp số/phanh; mã VFD, điện trở hãm, phối hợp MCCB và cáp. Vì vậy phần lựa chọn mã thiết bị, điện trở Ω/W và kiểm nghiệm phát nhiệt đầy đủ vẫn chưa được xem là hoàn thành. Báo cáo không tuyên bố cơ cấu đạt chứng nhận an toàn.')
sub('Tài liệu dự án')
for txt in ['[1] SYSTEM ARCHITECTURE – CRANE TROLLEY DRIVE SYSTEM.txt: cấu trúc chức năng và phân cấp phần cứng.', '[2] Tính toán thông số động cơ.xlsx: sheet a=1 là cơ sở cấu hình, sheet a=6 và Copy of Tính toán chi tiết động dùng đối chiếu x=6 và x=3. Các ô chính E5:E24, C34:C49, M5:M38.', '[3] SOW_HungNB.docx và ảnh yêu cầu do người dùng cung cấp: lực cản, đồ thị chu kỳ, motor, biến tần, bảo vệ và hãm. Người dùng xác nhận chuyển động ngang và x=2.']:
    p(txt)
sub('Tài liệu kỹ thuật tham khảo')
refs=[('[4] SEW EURODRIVE, Project Planning for Drives: mô hình lực cản và bám bánh.','https://download.sew-eurodrive.com/download/pdf/10522913.pdf'),('[5] ABB, ACS355: tài liệu tham khảo chức năng drive/STO; không phải mã VFD đã chốt.','https://www.abb.com/global/en/areas/motion/drives/low-voltage-ac-drives/legacy-drives/acs355'),('[6] IEC 60204-32:2023: thiết bị điện máy nâng.','https://webstore.iec.ch/en/publication/63094'),('[7] ISO 13850:2015: nguyên tắc chức năng dừng khẩn.','https://committee.iso.org/standard/59970.html?browse=ics')]
for label,url in refs:
    pp=p(label+'\n'+url)
    for rr in pp.runs:rr.font.size=Pt(10)
d.core_properties.title='Thiết kế hệ truyền động xe con cầu trục — x = 2 theo mẫu'
d.core_properties.subject='Tính toán phần cứng 5,5 kW, chu kỳ S-curve và logic trạng thái'
d.save(OUT/'Bao_cao_xe_con_x2_theo_mau.docx')
print('Saved revised x2 report.')
