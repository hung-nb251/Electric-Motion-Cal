from pathlib import Path
import sys,os
sys.path.insert(0,str(Path(__file__).resolve().parent/'pydeps'))
os.environ['MPLCONFIGDIR']=str(Path(__file__).resolve().parent/'mplconfig')
import math,json,re
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.patches import FancyBboxPatch, FancyArrowPatch
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parents[2]
OUT=ROOT/'outputs/bao_cao_xe_con'; TMP=ROOT/'tmp/report_build'
OUT.mkdir(parents=True,exist_ok=True)
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':10})
m=12444.; m0=222.; g=9.81; r=.125; i=100.; eta=.96; J=.015; v=.12; f=.01; k=1.2
Ta=2.; tj=.25; ap=v/(Ta-tj); jp=ap/tj; L=9.; tc=L/v-Ta; tm=tc+2*Ta; Tcy=2*tm+60
C=m*g*f*r/(i*eta); B=m*r/(i*eta)+J*i/r
C0=m0*g*f*r/(i*eta); B0=m0*r/(i*eta)+J*i/r
Trms=math.sqrt(((C*C+C0*C0)*tm+(B*B+B0*B0)*2*ap*ap*(Ta-4*tj/3))/Tcy)

def profile(t):
    if t<tj:return .5*jp*t*t,jp*t,jp
    if t<Ta-tj:return .5*ap*tj+ap*(t-tj),ap,0
    if t<Ta:
        u=Ta-t;return v-.5*jp*u*u,jp*u,-jp
    if t<Ta+tc:return v,0,0
    if t<tm:
        vv,aa,jj=profile(t-(Ta+tc));return v-vv,-aa,-jj
    return 0,0,0
t=np.linspace(0,tm,7701); vals=np.array([profile(x) for x in t]); vel,acc,jerk=vals.T
tor=C+B*acc; powr=tor*vel*i/r
assert abs(np.trapezoid(vel,t)-L)<1e-5
assert max(abs(acc))<=.5 and max(abs(jerk))<=2
assert abs(Trms-.989297081243038)<1e-9
fig,axs=plt.subplots(3,1,figsize=(8,4.6),sharex=True)
for ax,y,label in zip(axs,[vel,acc,tor],['v (m/s)','a (m/s²)','M động cơ (N·m)']):
    ax.plot(t,y,color='#1F4E78',lw=1.7); ax.set_ylabel(label); ax.grid(alpha=.25)
axs[-1].set_xlabel('Thời gian một lượt có tải (s)');fig.tight_layout();fig.savefig(TMP/'profile.png',dpi=190);plt.close(fig)

def box(ax,x,y,w,h,label,color='#eef3f8',size=10):
    ax.add_patch(FancyBboxPatch((x-w/2,y-h/2),w,h,boxstyle='round,pad=0.02,rounding_size=0.05',facecolor=color,edgecolor='#34495e',lw=1))
    ax.text(x,y,label,ha='center',va='center',fontsize=size)
def arrow(ax,a,b,label='',offset=(0,0),color='#34495e'):
    ax.add_patch(FancyArrowPatch(a,b,arrowstyle='-|>',mutation_scale=12,color=color,lw=1.1))
    if label:ax.text((a[0]+b[0])/2+offset[0],(a[1]+b[1])/2+offset[1],label,ha='center',va='center',fontsize=8.5,backgroundcolor='white')

fig,ax=plt.subplots(figsize=(8,4.2));ax.set(xlim=(0,10),ylim=(0,5.2));ax.axis('off')
for x,lab in [(1.3,'Nguồn 3 pha\nMCCB + cách ly'),(3.8,'VFD vector\nSTO hai kênh'),(6.3,'Động cơ + phanh\nHộp số i ≈ 100'),(8.8,'Hai bánh chủ động\nXe con trên ray')]:box(ax,x,3.7,2.2,1,lab)
for x in [2.4,4.9,7.4]:arrow(ax,(x,3.7),(x+.3,3.7))
box(ax,3.8,1.5,2.2,1,'PLC + HMI\nS-curve và liên khóa');box(ax,7.2,1.5,2.5,1,'Encoder, hành trình\nPhanh, nhiệt, tải')
arrow(ax,(5.9,1.5),(4.95,1.5));arrow(ax,(3.8,2),(3.8,3.15),'Lệnh',(0.4,0));arrow(ax,(6.3,3.15),(7.2,2),'Đo lường',(0.6,0))
box(ax,1.3,1.5,2.1,1,'Relay / PLC an toàn\nE-stop + cực hạn',color='#f7eceb',size=9)
arrow(ax,(1.3,2),(2.7,3.2),'STO / phanh',(-.1,.25),color='#a23b32')
fig.tight_layout();fig.savefig(TMP/'architecture.png',dpi=190);plt.close(fig)

fig,ax=plt.subplots(figsize=(8.2,8.3));ax.set(xlim=(0,10),ylim=(0,11));ax.axis('off')
states=[(9.9,'S0  KHÓA KHỞI ĐỘNG\nSTO tác động, phanh đóng'),(8.35,'S1  SẴN SÀNG\nChờ lệnh mới hợp lệ'),(6.8,'S2  TẠO MÔ-MEN\nEnable VFD, tốc độ đặt bằng 0'),(5.25,'S3  MỞ PHANH\nChờ xác nhận đã mở'),(3.7,'S4  CHẠY TIẾN / LÙI\nBám S-curve, giám sát liên tục'),(2.15,'S5  GIẢM TỐC\nDừng có điều khiển'),(.6,'S6  ĐÓNG PHANH\nXác nhận phanh rồi bỏ mô-men')]
for y,lab in states:box(ax,3.3,y,4.5,.95,lab,size=10)
labs=['Tự kiểm tra + Reset hợp lệ','Lệnh mới + điều kiện cho phép','DriveReady + TorqueReady','BrakeOpen xác nhận','Stop / nhả nút / chạm giới hạn','Tốc độ gần 0 ổn định']
for a,b,lab in zip(states[:-1],states[1:],labs):arrow(ax,(3.3,a[0]-.48),(3.3,b[0]+.48),lab,(.0,.0))
box(ax,8,7.5,3.3,1.3,'S7  LỖI ĐÃ CHỐT\nCấm chạy, lưu mã lỗi\nBảo vệ theo loại lỗi',color='#f7eceb',size=10)
box(ax,8,4.7,3.3,1.3,'S8  DỪNG KHẨN\nSTO + cắt điện nhả phanh\nMạch an toàn độc lập',color='#f7eceb',size=10)
arrow(ax,(5.6,3.7),(6.3,4.6),'E-stop',(.1,-.3),color='#a23b32')
arrow(ax,(5.6,6.8),(6.3,7.3),'Lỗi mất điều khiển',(.4,-.5),color='#a23b32')
arrow(ax,(8,5.37),(8,6.8),'Đã dừng + chốt lỗi',(.0,0),color='#a23b32')
ax.plot([1, .55,.55,1],[.6,.6,8.35,8.35],color='#34495e',lw=1)
arrow(ax,(.55,8.35),(1,8.35));ax.text(.3,4.3,'Phanh đóng xác nhận',rotation=90,ha='center',fontsize=8.5)
ax.plot([8,8,5.6],[8.17,9.9,9.9],color='#a23b32',lw=1)
arrow(ax,(6.1,9.9),(5.55,9.9),color='#a23b32');ax.text(7.5,10.35,'Hết lỗi + Reset tay\n+ mọi lệnh chạy đã nhả',ha='center',fontsize=8.5)
fig.tight_layout();fig.savefig(TMP/'states.png',dpi=190);fig.savefig(OUT/'Luu_do_trang_thai.svg');plt.close(fig)

d=Document(); sec=d.sections[0];sec.page_width=Inches(8.5);sec.page_height=Inches(11);sec.top_margin=Inches(.7);sec.bottom_margin=Inches(.65);sec.left_margin=Inches(.8);sec.right_margin=Inches(.7)
for nm in ['Normal','Body Text','Title','Subtitle','Heading 1','Heading 2','Caption']:
    s=d.styles[nm];s.font.name='Times New Roman';s.font.color.rgb=RGBColor(0,0,0);s.font.size=Pt(12)
    rf=s.element.get_or_add_rPr().find(qn('w:rFonts'))
    if rf is not None:
        for key in list(rf.attrib):
            if key.endswith('Theme'):del rf.attrib[key]
        rf.set(qn('w:ascii'),'Times New Roman');rf.set(qn('w:hAnsi'),'Times New Roman')
for sty in d.styles:
    for border in list(sty.element.iter(qn('w:pBdr'))):border.getparent().remove(border)
d.styles['Caption'].font.size=Pt(10)
d.styles['Normal'].paragraph_format.space_after=Pt(6);d.styles['Normal'].paragraph_format.line_spacing=1.08
for nm,sz in [('Title',21),('Heading 1',16),('Heading 2',13)]:d.styles[nm].font.size=Pt(sz);d.styles[nm].font.bold=True
footer=sec.footer.paragraphs[0];footer.alignment=2
run=footer.add_run('Trang ');run.font.size=Pt(10)
fld=OxmlElement('w:fldSimple');fld.set(qn('w:instr'),'PAGE');footer._p.append(fld)
def p(txt,style=None):return d.add_paragraph(txt,style)
_new_page=False
def h(txt):
    global _new_page
    pp=d.add_heading(txt,level=1)
    if _new_page:pp.paragraph_format.page_break_before=True
    _new_page=False
def sub(txt):d.add_heading(txt,level=2)
def eq(txt):
    pp=d.add_paragraph();pp.paragraph_format.space_after=Pt(7)
    om=OxmlElement('m:oMath')
    def mr(text):
        rr=OxmlElement('m:r');tt=OxmlElement('m:t');tt.text=text;tt.set(qn('xml:space'),'preserve');rr.append(tt);return rr
    if any(x in txt for x in ['CMD_VALID','DIR_OK','PERMIT',' AND ']):
        om.append(mr(txt))
    else:
        pattern=r'([A-Za-zαωη])_([A-Za-zÀ-ỹ]+(?:,[A-Za-zÀ-ỹ]+)?)'
        pos=0
        for mt in re.finditer(pattern,txt):
            if mt.start()>pos:om.append(mr(txt[pos:mt.start()]))
            subnode=OxmlElement('m:sSub');base=OxmlElement('m:e');base.append(mr(mt.group(1)));subpart=OxmlElement('m:sub');subpart.append(mr(mt.group(2)));subnode.append(base);subnode.append(subpart);om.append(subnode);pos=mt.end()
        if pos<len(txt):om.append(mr(txt[pos:]))
    pp._p.append(om)
def tab(headers,rows,widths=None):
    t=d.add_table(rows=1,cols=len(headers));t.autofit=False
    if widths:
        for c,w in zip(t.columns,widths):c.width=Inches(w)
    for c,txt in zip(t.rows[0].cells,headers):c.text=txt
    rep=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(rep)
    for row in rows:
        for c,txt in zip(t.add_row().cells,row):c.text=str(txt)
    for ri,row in enumerate(t.rows):
        pr=row._tr.get_or_add_trPr();pr.append(OxmlElement('w:cantSplit'))
        for ci,c in enumerate(row.cells):
            if widths:c.width=Inches(widths[ci])
            cp=c._tc.get_or_add_tcPr();mar=OxmlElement('w:tcMar')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:w'),'70');e.set(qn('w:type'),'dxa');mar.append(e)
            cp.append(mar); borders=OxmlElement('w:tcBorders')
            for edge in ['top','left','bottom','right']:
                e=OxmlElement('w:'+edge);e.set(qn('w:val'),'single');e.set(qn('w:sz'),'4');e.set(qn('w:color'),'D9D9D9');borders.append(e)
            cp.append(borders)
            if ri==0:
                sh=OxmlElement('w:shd');sh.set(qn('w:fill'),'DCE6F1');cp.append(sh)
            for pp in c.paragraphs:
                pp.paragraph_format.space_after=Pt(2);pp.paragraph_format.line_spacing=1
                for rr in pp.runs:rr.font.size=Pt(10.5);rr.bold=ri==0
    d.add_paragraph().paragraph_format.space_after=Pt(0)
    return t
def page():
    global _new_page
    _new_page=True
def pic(name,w=6.8):d.add_picture(str(TMP/name),width=Inches(w))

d.add_paragraph('Thiết kế phần cứng và logic an toàn cho cơ cấu xe con cầu trục',style='Title')
p('Tải trọng 12 222 kg    |    Tốc độ di chuyển ngang 0,12 m/s',style='Subtitle')
h('1 Phạm vi và kết quả thiết kế')
p('Báo cáo tính chọn truyền động di chuyển ngang của xe con: tải quy đổi, bánh xe, hộp giảm tốc, động cơ, phanh và thiết bị điện; đồng thời xây dựng lưu đồ trạng thái để vận hành và xử lý sự cố. Cấu trúc thuyết minh kế thừa hai cấp hệ thống trong tài liệu SYSTEM ARCHITECTURE và các nhóm tính toán của bảng Excel khóa trước.')
p('Phương án đề xuất có điều kiện: một động cơ không đồng bộ ba pha 2,2 kW, 6 cực, khoảng 960 vòng/phút; hộp giảm tốc i ≈ 100; bánh xe đường kính 250 mm; hai trong bốn bánh được dẫn động cơ khí. Kết quả dựa trên cản lăn quy đổi 0,01 và quán tính phần quay quy về trục động cơ không vượt 0,015 kg·m². Những giá trị này là giả thiết thiết kế, phải được thay bằng số liệu thiết bị thực trước khi chốt cấu hình.')
tab(['Đầu vào được giao','Ký hiệu','Giá trị'],[['Tải nâng định mức','Q','12 222 kg'],['Khối lượng cơ cấu nâng hạ','m_h','222 kg'],['Chiều cao nâng','h','12 m'],['Tốc độ xe con tối đa','v_max','0,12 m/s'],['Giới hạn gia tốc','a_lim','0,5 m/s²'],['Giới hạn độ giật','j_lim','2 m/s³']],[3.7,1,2.3])
p('Khối lượng di chuyển có tải tối thiểu là Q + m_h = 12 444 kg. Khối lượng không tải tối thiểu là 222 kg. Chưa có khối lượng khung xe, động cơ, hộp số và phụ kiện; tổng thực phải cộng thêm m_x. Giá trị m_x = 0 trong phép tính cơ sở chỉ là mốc tính theo dữ liệu đã giao, không khẳng định xe con không có tự trọng.')
p('Chiều cao h = 12 m không phải hành trình ngang. Báo cáo không tính tang, cáp và động cơ nâng hạ. Các giới hạn gia tốc và độ giật được áp dụng cho chuyển động vận hành có điều khiển; dừng do mất điện/phanh cơ phải được kiểm chứng riêng.')

page();h('2 Đối chiếu tài liệu mẫu và giả thiết')
tab(['Điểm cần sửa','Cách xử lý trong báo cáo'],[['Mẫu kiến trúc ghi bánh xe nhưng chọn tỷ số truyền theo tang','Chỉ quy đổi từ vận tốc xe con và bán kính bánh xe.'],['Excel dùng μ = 0,2 cho lực cản chạy','Tách hệ số cản lăn f_r khỏi hệ số bám μ_b; giữ phép tính theo 0,2 ở mục đối chiếu.'],['Mô-men đã chia hiệu suất nhưng công suất lại chia tiếp','Công suất cơ trên trục động cơ bằng M × ω; không chia η lần thứ hai.'],['Lấy công suất trung bình từng đoạn để tính phát nhiệt','Tính mô-men RMS theo thời gian và kiểm tra mô-men đỉnh.'],['Tăng tốc tuyến tính ghép đột ngột','Dùng S-curve có giới hạn jerk, gia tốc liên tục.'],['Các tab dùng tải, tỷ số truyền và động cơ khác nhau','Không lấy danh mục thiết bị 15 kW của tab khác làm cấu hình của đề tài này.']],[2.9,4.1])
tab(['Thông số bổ sung','Cơ sở / giả thiết tính'],[['D = 0,25 m; η = 0,96; k = 1,2','Kế thừa SOW_HungNB.docx; kiểm tra lại η bằng catalog hộp số.'],['i = 100; 4 bánh, 2 bánh chủ động','Đề xuất bố trí để tính sơ bộ; tải trọng phân bố đều.'],['f_r = 0,01; μ_b = 0,20','f_r là giả thiết cản chạy; μ_b là giả thiết bám, không phải cùng một đại lượng.'],['J_rot = 0,015 kg·m²','Ngân sách tổng quán tính phần quay quy về trục động cơ, gồm rotor và truyền động.'],['L = 9 m; nghỉ 30 s tại mỗi đầu','Chu kỳ minh họa kế thừa quãng đường xấp xỉ trong Excel; không phải yêu cầu đã giao.'],['Ray ngang, trong nhà','Bỏ qua gió và độ dốc ở ca cơ sở; phải bổ sung nếu điều kiện thực khác.']],[2.85,4.15])
p('Theo mô hình cản chạy của SEW [4], f_r = (2e + μ_L d_b)/D + c. Với ví dụ e = 0,5 mm, μ_L = 0,005, d_b = 60 mm, D = 250 mm và c = 0,003, hệ số suy ra là 0,0082. Báo cáo dùng 0,01 làm giả thiết tròn có dự trữ; đây không phải hệ số đã đo của xe con. Các số liệu ví dụ không thay cho kiểm tra ổ lăn, lệch ray và ma sát vành bánh thực tế.')

page();h('3 Tải trọng và lực kéo cơ cấu')
eq('m = Q + m_h + m_x;   m₀ = m_h + m_x')
eq('F_c = f_r m g;   F_k = F_c + m a')
p('Trong các công thức dưới đây, F_k là lực tiếp tuyến tổng tại các bánh chủ động khi chuyển động theo chiều dương trên ray ngang. Khi giảm tốc vẫn dùng gia tốc có dấu; không tự động coi mọi đoạn giảm tốc đều tái sinh.')
tab(['Đại lượng ở ca cơ sở','Công thức','Kết quả'],[['Trọng lượng toàn hệ','m g','122 075,64 N'],['Lực cản đều','0,01 m g','1 220,76 N'],['Lực quán tính ở a = 0,5','m a','6 222,00 N'],['Lực kéo bao trên','F_c + m a_lim','7 442,76 N'],['Lực bám khả dụng 2 bánh','μ_b × 0,5 m g','12 207,56 N'],['Hệ số dự trữ bám','F_bám / F_k','1,64'],['Tỷ lệ tải trên bánh chủ động tối thiểu','F_k / (μ_b m g)','30,49 %']],[2.65,2.3,2.05])
p('Điều kiện không trượt là F_k ≤ μ_b N_chủ_động. Giả thiết hai bánh chịu 50% tổng tải thỏa điều kiện ở ca có tải. Với tải phân bố lệch phải tính phản lực từng bánh và dùng tổng phản lực nhỏ nhất trên các bánh chủ động; không lấy số bánh chia đều khi chưa kiểm tra vị trí trọng tâm.')
eq('F_c,tổng = f_r m g cos(θ) + m g sin(θ) + F_gió + F_cáp')
p('Nếu ray dốc hoặc dây dẫn di động tạo thêm lực kéo, đưa các thành phần đó vào lực cản. Mỗi 1% độ dốc bổ sung gần 1 221 N tại tải cơ sở, gần bằng lực cản lăn đã giả định; vì vậy độ dốc không được bỏ qua khi nghiệm thu.')
sub('Kiểm tra trường hợp giữ hệ số 0,2 của mẫu')
eq('F_c,mẫu = 0,2 m g = 24 415,13 N')
eq('F_k,mẫu = m (0,2 g + 0,5) = 30 637,13 N')
p('Lực kéo này lớn hơn lực bám 12 207,56 N của hai bánh chủ động. Ngay cả bốn bánh cùng chủ động với μ_b = 0,2, lực bám cũng chỉ bằng lực cản đều và không còn dự trữ tăng tốc. Vì vậy kết quả công suất 4,60 kW theo mẫu chỉ dùng đối chiếu toán học; không đủ cơ sở chốt phương án bánh thép chạy trên ray.')

page();h('4 Thông số truyền động cơ khí')
eq('r = D/2 = 0,125 m;   ω_b = v/r = 0,96 rad/s')
eq('n_b = 60 v/(π D) = 9,167 vòng/phút')
eq('n_m = i n_b = 916,73 vòng/phút;   ω_m = 96 rad/s')
p('Chọn sơ bộ hộp giảm tốc bánh răng với i ≈ 100. Động cơ 6 cực có tốc độ danh định dự kiến 960 vòng/phút; tốc độ làm việc yêu cầu nằm dưới tốc độ cơ sở. Ước tính tỷ lệ tần số khoảng 47,75 Hz, nhưng tốc độ cuối cùng được đóng vòng theo encoder và bù trượt, không ấn định bằng tỷ lệ tần số đơn thuần.')
eq('M_b,max = F_k,max r = 930,34 N·m')
eq('M_b,thiết kế = k M_b,max = 1 116,41 N·m')
tab(['Hạng mục','Yêu cầu tính chọn sơ bộ'],[['Bộ bánh xe','4 bánh D = 250 mm; 2 bánh chủ động, có dẫn hướng phù hợp ray.'],['Tải tĩnh mỗi bánh','30,52 kN nếu tải đặt giữa và phân bố đều.'],['Tải thiết kế sơ bộ mỗi bánh','1,25 × 30,52 = 38,15 kN; chọn cấp tải tối thiểu 40 kN trong ca phân bố đều.'],['Hộp số và truyền động chung','Tỷ số khoảng 100; mô-men đầu ra tính toán ≥ 1,12 kN·m, kiểm tra thêm mô-men tối đa do VFD và phanh.'],['Nhánh dẫn động mỗi bánh','Nếu chia tải đều, mô-men thiết kế khoảng 558,21 N·m mỗi nhánh.'],['Trục chung theo xoắn đơn thuần','Với ứng suất cắt cho phép giả định 40 MPa, d_min ≈ 52,2 mm; dự kiến d = 60 mm để tiếp tục kiểm tra.']],[2.25,4.75])
eq('d_min = ∛[16 M_thiết kế/(π [τ])]')
p('Trong công thức đường kính, đổi mô-men sang N·mm. Đường kính 60 mm chưa phải kích thước chế tạo: phải kiểm tra uốn, xoắn phối hợp, tập trung ứng suất tại then, độ võng, ổ lăn, tải hướng kính và tải dọc trục. Bề rộng vành bánh, vật liệu, độ cứng và tiếp xúc bánh–ray cần dữ liệu ray và phân cấp làm việc; không suy ra chỉ từ đường kính 250 mm.')
p('Hệ số 1,25 phân bố/động lực và ứng suất 40 MPa là giả thiết sơ bộ của báo cáo, không được gán là giá trị bắt buộc của TCVN. Kết cấu khung và dầm cầu nằm ngoài phạm vi SOW.')

page();h('5 Quỹ đạo vận tốc có giới hạn độ giật')
p('Chọn thời gian tăng tốc và giảm tốc vận hành T_a = 2 s, thời gian tăng/giảm gia tốc t_j = 0,25 s. Giữa hai đoạn jerk là 1,50 s gia tốc không đổi. Bộ tạo quỹ đạo phải giữ liên tục vận tốc và gia tốc khi nhận Stop hoặc đổi yêu cầu.')
eq('a_p = v_max/(T_a − t_j) = 0,068571 m/s²')
eq('j_p = a_p/t_j = 0,274286 m/s³')
p('Hai giá trị đều nhỏ hơn giới hạn đề bài. Trong đoạn tăng tốc: 0 ≤ t < t_j dùng a = j_p t; t_j ≤ t < T_a − t_j dùng a = a_p; T_a − t_j ≤ t ≤ T_a dùng a = j_p (T_a − t). Tích phân gia tốc cho vận tốc; đoạn giảm tốc đối xứng.')
tab(['Đại lượng','Giá trị'],[['Quãng đường tăng tốc / giảm tốc','0,5 v_max T_a = 0,12 m cho mỗi đoạn'],['Hành trình minh họa','9 m ngang'],['Thời gian chạy đều','(9 − 2 × 0,12)/0,12 = 73 s'],['Thời gian một lượt','2 + 73 + 2 = 77 s'],['Chu kỳ có tải đi và không tải về','77 + 30 + 77 + 30 = 214 s'],['Tỷ lệ thời gian chuyển động','154/214 = 71,96 %']],[3.9,3.1])
pic('profile.png',5.7)
p('Hình 1. Vận tốc, gia tốc và mô-men của lượt có tải theo ca cơ sở. Mô-men trục động cơ đã gồm quán tính phần quay giả định.',style='Caption')
p('Quỹ đạo tăng tốc nhanh nhất từ trạng thái nghỉ với j = 2 có a_đỉnh = √(v_max j) = 0,48990 m/s² và T_a,min = 2√(v_max/j) = 0,48990 s. Không thể áp đồng thời a = 0,5 trong cả thời gian v/a và giữ jerk hữu hạn. Kiểm tra a = 0,5 ở mục động cơ là bao trên bảo thủ, không phải quỹ đạo thực đồng thời với v_max.')

page();h('6 Quy đổi mô-men và chọn động cơ')
eq('J_tịnh tiến = m (r/i)² = 0,01944375 kg·m²')
eq('α_m = (i/r) a;   M_m = (F_c + m a) r/(η i) + J_rot α_m')
eq('P_m = M_m ω_m;   P_điện ≈ P_m/(η_motor η_VFD)')
p('Mô-men là tổng mô-men kéo tải đã quy đổi qua hộp số và mô-men tăng tốc phần quay. Không cộng J_tịnh tiến α_m lần nữa nếu đã có m a trong lực kéo. Khi dòng công suất qua hộp số đảo chiều, dùng hiệu suất truyền ngược; ở quỹ đạo vận hành đang xét, lực tại bánh vẫn cùng chiều vận tốc, nên công thức trên áp dụng được cả khi rotor giảm tốc gây mô-men trục âm.')
tab(['Trường hợp có tải','Kết quả'],[['Mô-men chạy đều','1,5895 N·m'],['Mô-men đỉnh quỹ đạo T_a = 2 s','3,5235 N·m'],['Mô-men bao trên tại a = 0,5','15,6911 N·m'],['Mô-men thiết kế có k = 1,2','18,8293 N·m'],['Công suất bao trên có dự trữ tại ω = 96','1,8076 kW'],['Động cơ đề xuất 2,2 kW, 960 vòng/phút','M_đm = 21,8838 N·m'],['Tỷ số M_đm / M_thiết kế','1,162']],[4.4,2.6])
p('Công suất 1,8076 kW là bao trên lấy mô-men và tốc độ cực đại cùng lúc, không phải đỉnh chính xác của S-curve. Động cơ 2,2 kW đáp ứng bao trên theo giả thiết. Chọn động cơ dùng được với biến tần, có cảm biến nhiệt, encoder và phanh lò xo; kiểm tra đặc tính mô-men liên tục tại khoảng 917 vòng/phút và chế độ làm việc thực.')
sub('Kiểm tra phát nhiệt theo chu kỳ minh họa')
eq('M_RMS = √[ ∫ M_m²(t) dt / T_chu kỳ ] = 0,9893 N·m')
p('Tích phân gồm lượt có tải và không tải, mỗi lượt 77 s, nghỉ 30 s mỗi đầu với phanh giữ và động cơ bỏ mô-men. Đặt C = f_r m g r/(η i), B = m r/(η i) + J_rot i/r. Tích phân của mỗi lượt là C² T_lượt + 2 B² a_p² (T_a − 4t_j/3); cộng hai lượt rồi chia 214 s. Không thay bằng RMS của công suất trung bình từng đoạn.')
p('M_RMS thấp hơn nhiều mô-men định mức, nhưng đây là chỉ số tải cơ tương đương; chưa thay cho kiểm tra nhiệt bằng catalog vì còn tổn hao không tải, số lần khởi động, tự làm mát và thời gian giữ mô-men. Nhãn chế độ S1/S3/S4 chỉ được chốt khi có chu kỳ và dữ liệu nhà sản xuất.')

page();h('7 Độ nhạy và đối chiếu công suất theo mẫu')
tab(['Mô hình / thay đổi','Mô-men có dự trữ','Công suất bao trên'],[['Cơ sở f_r = 0,01; m_x = 0; J = 0,015','18,83 N·m','1,808 kW'],['Thêm khung và phụ kiện m_x = 500 kg','19,30 N·m','1,852 kW'],['Thêm khung và phụ kiện m_x = 1 000 kg','19,76 N·m','1,897 kW'],['f_r tăng lên 0,02; các biến khác giữ nguyên','20,74 N·m','1,991 kW'],['J_rot tăng lên 0,025 kg·m²','23,63 N·m','2,268 kW']],[3.8,1.6,1.6])
p('Bảng trên thay từng yếu tố riêng lẻ. Khi nhiều yếu tố cùng tăng, phải tính lại tổng hợp. Trường hợp J_rot = 0,025 không còn thỏa mô-men liên tục của động cơ 2,2 kW đã giả định; cần chọn lại động cơ hoặc giảm gia tốc đặt, đồng thời kiểm tra khả năng quá tải VFD.')
sub('Nếu giảng viên yêu cầu giữ nguyên mô hình hệ số 0,2')
eq('M_b = 30 637,128 × 0,125 = 3 829,641 N·m')
eq('M_m,k = 1,2 M_b/(0,96 × 100) = 47,8705 N·m')
eq('P_m,k = M_m,k × 96 = 4 595,57 W')
p('Theo đúng mô hình giản lược của bảng mẫu, công suất tiêu chuẩn kế tiếp là 5,5 kW. Con số 4,5956 kW chỉ tính khối lượng tịnh tiến và hiệu suất cơ khí một lần. Không cộng thêm một lần chia 0,96 vào công suất sau khi mô-men đã được quy đổi.')
p('Nếu giữ thêm quán tính rotor 0,039 kg·m² ghi trong tab mẫu a=1, mô-men tăng tốc rotor tại a = 0,5 là 15,6 N·m. Mô-men tổng có dự trữ trở thành 66,59 N·m, lớn hơn mô-men danh định khoảng 53,98 N·m của động cơ 5,5 kW ở 973 vòng/phút. Khi đó phải kiểm tra quá tải khoảng 1,24 lần, thời gian quá tải và quán tính hộp số; không thể kết luận chỉ từ điều kiện 5,5 > 4,60 kW.')
p('Lựa chọn chính của báo cáo là mô hình cản lăn và động cơ 2,2 kW có điều kiện. Phương án 5,5 kW chỉ là kết quả đối chiếu mô hình mẫu, do kiểm tra bám ở mục 3 không đạt. Không dùng hai công suất này như hai lựa chọn tương đương đã được xác nhận.')
sub('Giới hạn mô-men tại cơ cấu')
p('Động cơ 2,2 kW có thể tạo mô-men lớn hơn nhu cầu kéo tính toán. Phải cấu hình giới hạn mô-men theo tải, gia tốc và sức bền hộp số; đồng thời kiểm tra mô-men cực đại VFD khi kẹt bánh. Hộp số ≥ 1,12 kN·m chỉ đáp ứng tải thiết kế đã tính, chưa mặc nhiên chịu được toàn bộ mô-men quá tải cực đại của động cơ.')

page();h('8 Phanh và khoảng cách dừng')
eq('E_k = 0,5 m v² + 0,5 J_rot ω_m² = 158,72 J')
p('Đây là năng lượng động học ca cơ sở tại tốc độ cực đại, chưa gồm năng lượng tải lắc hoặc tác động ngoại lực. Dùng E_k làm mức trên sơ bộ cho năng lượng phải tiêu tán mỗi lần dừng trên ray ngang, khi bỏ qua cản lăn. Chọn phần tử hãm có khả năng chịu ít nhất 1,2 E_k = 190,46 J mỗi xung; vẫn phải kiểm tra công suất xung và chu kỳ lặp.')
sub('Điện trở hãm cho VFD')
eq('R ≥ R_min,VFD;   P_xung = U_chopper²/R;   E_xung = ∫ P_R(t) dt')
p('Không sao chép điện trở 200 Ω, 1 kW của mẫu vì chưa biết ngưỡng chopper, điện trở nhỏ nhất cho phép và năng lượng chấp nhận của VFD. Với chu kỳ minh họa, mức trên năng lượng hai lần dừng là 229,44 J/214 s ≈ 1,07 W trung bình; công suất trung bình nhỏ không đủ để chọn điện trở. Ca dừng nhanh, tải lắc và tần suất thao tác mới quyết định công suất xung.')
sub('Phanh cơ và mất điện')
p('Chọn cơ cấu phanh lò xo đóng khi mất điện, nhả bằng điện và có phản hồi trạng thái. Phanh phải được nhà sản xuất cho phép hãm động khi dùng làm phương tiện dừng khẩn; phanh chỉ dùng giữ tĩnh không được mặc nhiên sử dụng để dừng khẩn nhiều lần. Mô-men giữ tối thiểu tính từ độ dốc, gió và lực kéo cáp, với hiệu suất truyền ngược được kiểm chứng.')
p('Giá trị khảo sát ban đầu 5 N·m trên trục động cơ, giả định truyền ngược không tổn hao và đóng tức thời: khối lượng quán tính tương đương có tải 22 044 kg; lực hãm tương đương 4 000 N. Gia tốc chậm dần có tải khoảng 0,237 m/s² và không tải khoảng 0,409 m/s² nếu cộng cản chạy. Đây chỉ là ước lượng kiểm tra; chưa chốt mô-men phanh vì thời gian đóng, biến thiên mô-men, bám ray và hiệu suất truyền ngược chưa biết. Jerk của phanh cơ không được bảo đảm bởi S-curve của PLC.')
sub('Vị trí công tắc giảm tốc và cực hạn')
eq('s_dừng,thường = v_max T_d/2 = 0,12 m')
eq('d_kích hoạt ≥ v_max t_trễ + s_dừng + Δs')
p('Ví dụ t_trễ tổng 0,10 s và dự trữ Δs = 0,05 m cho d ≥ 0,182 m; đặt sơ bộ công tắc dừng vận hành cách điểm dừng mong muốn ít nhất 0,20 m. Không dùng 0,20 m làm khoảng cách an toàn đã chứng nhận. Cực hạn độc lập phải bố trí sao cho quãng đường dừng khẩn đo được vẫn nằm trước chặn cơ khí, kể cả trường hợp công tắc vận hành hỏng.')

page();h('9 Kiến trúc phần cứng hai cấp')
pic('architecture.png',5.6)
p('Hình 2. Luồng năng lượng, điều khiển và nhánh an toàn độc lập. Sơ đồ khối không thay cho bản vẽ đấu dây.',style='Caption')
tab(['Cấp 1','Cấp 2 và tiêu chí chọn'],[['Nguồn và bảo vệ','3 pha 380/400 V, 50 Hz; dao cách ly khóa được, MCCB/cầu chì theo bảng phối hợp VFD, PE, nguồn điều khiển 24 VDC.'],['Truyền động','Motor 2,2 kW 6 cực; encoder; phanh; hộp số i ≈ 100; hai bánh dẫn động.'],['Biến tần','Cấp 3 kW đáp ứng quy tắc SOW 1,3 × 2,2 = 2,86 kW; chọn theo dòng HD thực, vector có encoder và STO hai kênh.'],['Điều khiển','PLC sinh S-curve và logic trạng thái; HMI hiển thị trạng thái, tốc độ, tải và mã lỗi; VFD thực hiện vòng dòng và tốc độ.'],['Cảm biến','Encoder A/B/Z tương thích VFD; công tắc vận hành hai đầu; cực hạn độc lập; nhiệt motor; load-cell/tín hiệu quá tải từ bộ nâng; phản hồi phanh.'],['An toàn','Relay/PLC an toàn; E-stop hai kênh; giám sát reset và tiếp điểm; ngắt STO và điện nhả phanh khi cần.']],[1.6,5.4])
p('Dòng động cơ ước tính để lập dự toán: I ≈ 2 200/(√3 × 400 × 0,85 × 0,75) ≈ 4,98 A, trong đó hiệu suất 0,85 và cosφ = 0,75 là giả định. Chọn VFD có dòng liên tục sau giảm định mức ít nhất bằng dòng nhãn motor, kiểm tra mô-men/dòng quá tải và tương thích encoder. Không dùng dòng khởi động trực tiếp từ lưới làm dòng khởi động VFD.')
p('MCCB và tiết diện cáp chưa chốt mã: cần dòng đầu vào VFD, chiều dài cáp, cách lắp đặt, nhiệt độ, sụt áp và dòng ngắn mạch tại tủ. Điều kiện cơ bản I_tải ≤ I_bảo vệ ≤ I_cho phép,cáp; khả năng cắt phải lớn hơn dòng ngắn mạch dự kiến. Mạch motor do VFD cấp còn phải phối hợp bảo vệ nhiệt và hướng dẫn hãng.')
p('Dùng cảm biến dòng nội bộ VFD cho điều khiển vector; cảm biến ngoài nếu cần phải có băng thông và cách ly phù hợp. CT 50 Hz kiểu SCT013 trong mẫu không được coi là cảm biến phản hồi trực tiếp thích hợp cho vòng dòng PWM. ACS355 là mã tham khảo lịch sử; không khẳng định cấu hình cụ thể hoặc phụ kiện đã được lựa chọn cho phương án này [5].')

page();h('10 Tín hiệu và các điều kiện liên khóa')
tab(['Nhóm','Tín hiệu tối thiểu','Ý nghĩa'],[['Lệnh người vận hành','FWD, REV, STOP, RESET','Lệnh tiến/lùi kiểu giữ để chạy; Reset cạnh lên riêng.'],['An toàn độc lập','ESTOP_CH1/2, FINAL_L/R, EDM','E-stop, cực hạn hai đầu và giám sát phần tử chấp hành.'],['Hành trình vận hành','LIMIT_L_OK, LIMIT_R_OK','Cho phép chạy theo hướng; tiếp điểm NC, kiểm tra đứt dây.'],['Biến tần','DRIVE_READY, FAULT, TORQUE_READY','Điều kiện phát mô-men; lỗi mất điều khiển phải phản ứng ngay.'],['Cơ khí và tải','BRAKE_OPEN, BRAKE_CLOSED, LOAD_OK','Trạng thái phanh có kiểm tra hợp lý; tải trong ngưỡng được hiệu chuẩn.'],['Giám sát','SPEED_VALID, ZERO_SPEED, TEMP_OK','Encoder hợp lệ; tốc độ gần 0 đủ thời gian; nhiệt trong giới hạn.'],['Đầu ra điều khiển','ENABLE, SPEED_REF, BRAKE_RELEASE','Lệnh drive, tốc độ và nhả phanh; nhánh an toàn có quyền cắt độc lập.'],['Đầu ra thông báo','FAULT_LATCH, BUZZER, STATUS','Mã lỗi lưu giữ và cảnh báo trạng thái.']],[1.35,2.7,2.95])
eq('CMD_VALID = FWD XOR REV')
eq('DIR_OK = (FWD AND LIMIT_R_OK) OR (REV AND LIMIT_L_OK)')
eq('PERMIT = SAFETY_OK AND DRIVE_READY AND TEMP_OK')
eq('             AND LOAD_OK AND SPEED_VALID AND NOT FAULT_LATCH')
p('Ngoài các biểu thức trên, khởi động yêu cầu phanh đang đóng, các tín hiệu phanh không mâu thuẫn, lệnh Stop không tác động và bộ điều khiển đã ghi nhận tất cả nút chạy được nhả kể từ lần reset. FWD và REV cùng tác động phải đưa về dừng, không tự chọn ưu tiên một chiều.')
p('Chạm giới hạn vận hành bên phải chỉ cấm chạy tiếp sang phải; có thể rút về trái bằng lệnh mới sau khi dừng và kiểm tra điều kiện. Nếu đứt dây hoặc tín hiệu không hợp lý, chốt lỗi và không cho rút tự động. Chạm cực hạn an toàn khóa cả hai chiều; chỉ phục hồi theo quy trình bảo trì, không bypass bằng HMI.')
p('Giới hạn tải của cơ cấu nâng phải lấy từ hệ cân tải đã hiệu chuẩn. Báo cáo không tự đặt ngưỡng 110% hoặc 125%. Khi quá tải trong lúc di chuyển, dừng có kiểm soát nếu drive còn điều khiển được và khóa chạy mới; trình tự xử lý tải thuộc quy trình tổng thể cầu trục.')
p('ZERO_SPEED khởi tạo kiểm chứng ở |v| < 0,002 m/s liên tục 0,2 s; đây là giá trị cài đặt thử. Bộ giám sát an toàn cần tín hiệu/thiết bị có cấp an toàn phù hợp nếu trạng thái đứng yên được dùng cho chức năng bảo vệ con người. PLC thường và encoder thường không tự tạo thành chức năng an toàn đã chứng nhận.')

page();h('11 Lưu đồ trạng thái vận hành')
pic('states.png',6.6)
p('Hình 3. Luồng chính S0 đến S6, nhánh lỗi S7 và dừng khẩn S8. Nhánh an toàn có hiệu lực từ mọi trạng thái; các mũi tên lỗi trong hình chỉ minh họa vị trí kết nối. Lỗi còn điều khiển được đi qua S5 và S6 trước khi vào S7.',style='Caption')

page();h('12 Bảng chuyển trạng thái và xử lý lỗi')
tab(['Trạng thái','Điều kiện chuyển tiếp / tác động'],[['S0 Khóa khởi động','STO tác động, cấm nhả phanh. Tự kiểm tra hợp lệ + Reset tay + đã nhả lệnh chạy → S1.'],['S1 Sẵn sàng','Phanh đóng, mô-men bằng 0. Lệnh mới hợp lệ + PERMIT + DIR_OK → S2; lệnh còn giữ từ trước không có hiệu lực.'],['S2 Tạo mô-men','Bỏ STO qua mạch an toàn khi được phép, Enable VFD ở tốc độ 0; DriveReady + TorqueReady → S3. Quá thời gian → S7.'],['S3 Mở phanh','Ra lệnh nhả phanh, giữ tốc độ 0. Xác nhận BrakeOpen → S4. Mất cho phép → hủy khởi động và đóng phanh; lỗi/timeout → S7.'],['S4 Chạy','Bám tốc độ đặt có giới hạn a và j. Stop, nhả nút, đổi chiều, giới hạn vận hành hoặc lỗi còn điều khiển được → S5.'],['S5 Giảm tốc','Đưa tốc độ về 0 bằng quỹ đạo liên tục. Giữ phanh mở trong khi drive còn điều khiển. ZeroSpeed ổn định → S6. Timeout hoặc mất điều khiển → nhánh dừng bảo vệ.'],['S6 Đóng phanh','Cắt lệnh nhả phanh; giữ mô-men tại 0 tốc độ đến xác nhận BrakeClosed. Sau đó bỏ mô-men / STO → S1 hoặc S7 nếu đang chờ chốt lỗi.'],['S7 Lỗi đã chốt','Cấm chạy. Reset chỉ được chấp nhận khi nguyên nhân đã hết, an toàn hợp lệ, phanh đã đóng, xe đứng yên và các nút chạy đã nhả → S0.'],['S8 Dừng khẩn','E-stop/cực hạn/mạch an toàn mất hợp lệ → STO + mất điện nhả phanh bằng mạch độc lập; chốt lỗi. Nhả E-stop không tự chạy lại.']],[1.35,5.65])
sub('Phân loại phản ứng lỗi')
p('Lỗi còn kiểm soát được: cảnh báo nhiệt, quá tải tín hiệu, mất lệnh truyền thông nhưng VFD còn khả năng dừng, giới hạn vận hành. Yêu cầu dừng có điều khiển; sau khi phanh đóng mới bỏ mô-men và chốt lỗi khi cần. Không mặc định mọi quá áp/quá dòng đều thuộc nhóm này.')
p('Lỗi mất điều khiển: VFD trip, quá dòng cứng, mất encoder làm mất kiểm soát, mất nguồn, mạch an toàn tác động. Không chờ S-curve hoặc TorqueReady. Nhánh độc lập cắt STO và nguồn nhả phanh; phanh và chặn cơ khí phải đủ khả năng cho quãng đường dừng thực. Nếu BrakeClosed không xác nhận trong dừng thường, giữ mô-men khi drive còn khả năng, chốt lỗi và cấm rời trạng thái bảo vệ; không bỏ mô-men một cách mù quáng.')
p('Thiết kế này minh họa dừng khẩn loại 0 cho xe con ngang bằng STO kết hợp phanh hãm động. Nếu đánh giá rủi ro yêu cầu dừng loại 1, phải dùng SS1/timer an toàn được thẩm định rồi STO; không thay bằng timer PLC thường. STO chỉ ngăn tạo mô-men, không phải ngắt cách ly nguồn hoặc tự hãm cơ học [5–7].')

page();h('13 Quy trình vận hành và kiểm chứng')
sub('Trình tự vận hành')
for txt in ['1. Trước cấp điện: kiểm tra ray, bánh xe, vật cản, cáp cấp nguồn, phanh, cảm biến và vùng di chuyển; xác nhận tải không vượt định mức, người không đứng dưới tải.', '2. Cấp nguồn: mạch ở S0, phanh đóng. Kiểm tra E-stop và cực hạn hợp lệ, lệnh tiến/lùi đã nhả; nhấn Reset để sang S1.', '3. Chạy: giữ nút tiến hoặc lùi. Drive tạo mô-men tại 0 tốc độ, xác nhận rồi mới nhả phanh; chạy sau khi nhận BrakeOpen. Giám sát tải và trạng thái.', '4. Dừng/đổi chiều: nhả nút hoặc Stop; giảm tốc về 0, đóng phanh rồi bỏ mô-men. Đổi chiều chỉ sau khi dừng hoàn toàn và nhận lệnh mới.', '5. Sự cố: dùng E-stop khi cần dừng khẩn. Sau dừng, xử lý nguyên nhân và kiểm tra phanh; nhả E-stop, Reset rồi ra lệnh chạy riêng. Không nối tắt liên khóa.', '6. Kết thúc/bảo trì: đưa tải về trạng thái an toàn theo quy trình toàn cầu trục; đóng phanh, ngắt dao cách ly và khóa nguồn. Chờ xả DC bus và xác nhận hết điện theo hướng dẫn VFD.']:
    p(txt)
sub('Ma trận kiểm chứng trước nghiệm thu')
tab(['Phép kiểm tra','Kết quả cần đạt'],[['Chạy có tải / không tải','Đo |v| ≤ 0,12; |a| ≤ 0,5; |j| ≤ 2 trong chuyển động thường; công bố cách lọc/ước lượng đạo hàm.'],['Bám ray và phân bố tải','Không trượt trong điều kiện ray thực; phản lực bánh chủ động đủ ở tải lệch bất lợi.'],['Stop khi đang tăng tốc','Gia tốc không nhảy bậc; quỹ đạo dừng giữ jerk trong giới hạn.'],['Hai nút chiều cùng tác động','Không khởi động; nếu đang chạy thì dừng có điều khiển.'],['Hành trình vận hành / cực hạn','Chỉ rút khỏi giới hạn vận hành theo hướng cho phép; cực hạn khóa cả hai chiều.'],['Phanh không mở / không đóng','Timeout chốt lỗi; không chạy khi chưa mở; giữ mô-men khi đóng thất bại và drive còn khả năng.'],['E-stop, mất nguồn, VFD trip','Mạch an toàn tác động độc lập PLC; đo thời gian và quãng đường dừng trước chặn cơ khí.'],['Khôi phục điện và nhả E-stop','Không tự khởi động lại; cần Reset và lệnh chạy mới.'],['Nhiệt, hãm và chu kỳ lặp','Dòng, nhiệt motor/VFD/phanh và năng lượng điện trở trong định mức thực.']],[2.4,4.6])

page();h('14 Điều kiện chốt thiết kế và tài liệu tham khảo')
p('Kết quả tính chọn 2,2 kW là phương án thiết kế sơ bộ có điều kiện. Hồ sơ chế tạo cần bổ sung: m_x thực; bố trí trọng tâm và phản lực bánh; loại ray/bánh/ổ; độ dốc và lực kéo cáp; chu kỳ làm việc; catalog motor với J, M, I, hiệu suất và chế độ nhiệt; hiệu suất thuận/ngược hộp số; dữ liệu phanh; catalog VFD và chopper; dòng ngắn mạch, chiều dài và cách đi cáp. Từ đó cập nhật bảng tính và kiểm tra lại đồng thời lực bám, mô-men, nhiệt, tải cơ khí và dừng khẩn.')
p('Chiều cao nâng 12 m chỉ gợi ý cần xét dao động tải trong bài toán tổng thể. Chiều dài con lắc thực là khoảng treo tức thời, không luôn bằng chiều cao nâng. Nếu giả sử chiều dài treo 12 m, chu kỳ dao động nhỏ khoảng 6,95 s; S-curve làm chuyển động êm hơn nhưng không tự bảo đảm triệt lắc.')
p('SOW dẫn TCVN 4244:2005. Báo cáo không gán giới hạn a và j của đề bài cho một điều khoản tiêu chuẩn, cũng không tuyên bố hệ thống đã đạt chứng nhận an toàn. Việc chọn PLr/SIL, xác nhận mạch an toàn và xác định bộ tiêu chuẩn áp dụng phải thực hiện theo đánh giá rủi ro của thiết bị thực.')
sub('Tài liệu dự án')
for txt in ['[1] SYSTEM ARCHITECTURE – CRANE TROLLEY DRIVE SYSTEM.txt, mục Level 1 và Level 2. Dùng tham khảo cấu trúc chức năng; không kế thừa nguyên trạng định mức thiết bị.', '[2] Tính toán thông số động cơ.xlsx, các tab Tính toán chi tiết động cơ a=1, a=6 và Tính chọn thiết bị đi kèm. Đối chiếu các ô E6:E24, M5:M38 và thông số motor C34:C49 ở tab a=1.', '[3] SOW_HungNB.docx, mục phạm vi, bảng đầu vào và tiêu chí nghiệm thu. Người dùng xác nhận v, a và j áp dụng cho chuyển động ngang.']:
    p(txt)
sub('Tài liệu kỹ thuật tra cứu ngày 23 tháng 9 năm 2026')
refs=[('[4] SEW EURODRIVE, Project Planning for Drives, 2001, các mục 7.4 và 8 về lực cản, công suất và bám bánh.','https://download.sew-eurodrive.com/download/pdf/10522913.pdf'),('[5] ABB, trang sản phẩm ACS355 và thư viện hướng dẫn. Dùng tham khảo vector, STO và phanh; chưa chốt mã VFD.','https://www.abb.com/global/en/areas/motion/drives/low-voltage-ac-drives/legacy-drives/acs355'),('[6] IEC 60204-32:2023, phạm vi thiết bị điện của máy nâng.','https://webstore.iec.ch/en/publication/63094'),('[7] ISO 13850:2015, nguyên tắc chức năng dừng khẩn.','https://committee.iso.org/standard/59970.html?browse=ics'),('[8] Siemens, SINAMICS S210 Safety Integrated, 07/2025, mục SBC và giới hạn hãm khẩn của phanh. Tham khảo nguyên lý, không phải lựa chọn S210 cho đề tài.','https://cache.industry.siemens.com/dl/files/346/109994346/att_1340706/v1/S210_MC_SI_commiss_man_0725_en-US.pdf')]
for label,url in refs:
    pp=p(label+'\n'+url)
    for rr in pp.runs:rr.font.size=Pt(10)

d.core_properties.title='Thiết kế phần cứng và logic an toàn cho cơ cấu xe con cầu trục'
d.core_properties.subject='Tải 12 222 kg và tốc độ ngang 0,12 m/s'
d.save(OUT/'Bao_cao_phan_cung_va_logic_xe_con.docx')
(TMP/'results.json').write_text(json.dumps({'mass':m,'rolling_force':m*g*f,'torque_design':1.2*(C+B*.5),'rms':Trms,'normal_peak':C+B*ap,'normal_peak_power':float(max(powr)),'distance_integrated':float(np.trapezoid(vel,t)),'cycle':Tcy},indent=2),encoding='utf-8')
print((TMP/'results.json').read_text())
