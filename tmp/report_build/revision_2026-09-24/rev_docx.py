"""Sua bao cao Word ban nguoi dung (24/09/2026).

Muc 1-6: sua tai cho (giu dinh dang nguoi dung). Muc 7 tro di: dung lai.
Chay: python rev_docx.py <docx_vao> <thu_muc_hinh> <docx_ra>
"""
import copy
import re
import sys
from pathlib import Path

import docx
from docx.shared import Inches
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

SRC, FIG, DST = Path(sys.argv[1]), Path(sys.argv[2]), Path(sys.argv[3])
d = docx.Document(str(SRC))
body = d.element.body
W = "{http://schemas.openxmlformats.org/wordprocessingml/2006/main}"


def ptext(el):
    return "".join(n.text or "" for n in el.iter() if n.tag in (qn("w:t"), qn("m:t")))


def find_p(prefix):
    for el in body.iterchildren():
        if el.tag == qn("w:p") and ptext(el).strip().startswith(prefix):
            return el
    raise KeyError(prefix)


def tables():
    return [el for el in body.iterchildren() if el.tag == qn("w:tbl")]


def set_runs_text(p_el, text):
    """Giu rPr cua run dau, xoa cac run con lai (khong dung cho doan co OMML)."""
    runs = p_el.findall(qn("w:r"))
    first = runs[0]
    for r in runs[1:]:
        p_el.remove(r)
    for child in list(first):
        if child.tag != qn("w:rPr"):
            first.remove(child)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve")
    first.append(t)


# ---------------- mau dinh dang lay tu ban nguoi dung ----------------
T_H1 = copy.deepcopy(find_p("3 Chu kỳ phụ tải"))
T_H2 = copy.deepcopy(find_p("Quy đổi quán tính và bổ sung"))
T_P = copy.deepcopy(find_p("P_hd phục vụ tiêu chí"))
T_CAP = copy.deepcopy(find_p("Hình 1."))
T_EQ = copy.deepcopy(find_p("CMD_VALID"))
for tpl in (T_H1, T_H2, T_P, T_CAP):
    for br in tpl.iter(qn("w:lastRenderedPageBreak")):
        br.getparent().remove(br)
EQ_RPR = copy.deepcopy(T_EQ.find(".//" + qn("m:r")).find(qn("w:rPr")))
RUN_RPR = copy.deepcopy(T_P.find(qn("w:r")).find(qn("w:rPr")))


def mk(tpl, text):
    el = copy.deepcopy(tpl)
    set_runs_text(el, text)
    return el


def mr(text):
    r = OxmlElement("m:r"); r.append(copy.deepcopy(EQ_RPR))
    t = OxmlElement("m:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    return r


def mk_eq(text):
    el = copy.deepcopy(T_EQ)
    om = el.find(".//" + qn("m:oMath"))
    for c in list(om):
        om.remove(c)
    if any(k in text for k in ("PERMIT", "DIR_OK", " AND ", "CMD_")):
        om.append(mr(text)); return el
    pos = 0
    for m in re.finditer(r"([A-Za-zαωηψθ])_([A-Za-zÀ-ỹ0-9]+(?:,[A-Za-zÀ-ỹ]+)?)", text):
        if m.start() > pos:
            om.append(mr(text[pos:m.start()]))
        s = OxmlElement("m:sSub"); e = OxmlElement("m:e"); e.append(mr(m.group(1)))
        sb = OxmlElement("m:sub"); sb.append(mr(m.group(2))); s.append(e); s.append(sb); om.append(s)
        pos = m.end()
    if pos < len(text):
        om.append(mr(text[pos:]))
    return el


def sect():
    return body.find(qn("w:sectPr"))


def add(el):
    sect().addprevious(el)
    return el


def H1(t): return add(mk(T_H1, t))
def H2(t): return add(mk(T_H2, t))
def P(t): return add(mk(T_P, t))
def CAP(t): return add(mk(T_CAP, t))
def EQ(t): return add(mk_eq(t))


def PIC(name, width=6.7):
    d.add_picture(str(FIG / name), width=Inches(width))
    p_el = sect().getprevious()
    ppr = p_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr"); p_el.insert(0, ppr)
    jc = OxmlElement("w:jc"); jc.set(qn("w:val"), "center"); ppr.append(jc)
    keep = OxmlElement("w:keepNext"); ppr.insert(0, keep)
    return p_el


def cell_fmt(cell, header, width):
    tcpr = cell._tc.get_or_add_tcPr()
    tcw = tcpr.find(qn("w:tcW"))
    if tcw is None:
        tcw = OxmlElement("w:tcW"); tcpr.append(tcw)
    tcw.set(qn("w:w"), str(int(width * 1440))); tcw.set(qn("w:type"), "dxa")
    borders = OxmlElement("w:tcBorders")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + edge); e.set(qn("w:val"), "single"); e.set(qn("w:sz"), "4")
        e.set(qn("w:space"), "0"); e.set(qn("w:color"), "D9D9D9"); borders.append(e)
    tcpr.append(borders)
    if header:
        sh = OxmlElement("w:shd"); sh.set(qn("w:val"), "clear"); sh.set(qn("w:color"), "auto")
        sh.set(qn("w:fill"), "DCE6F1"); tcpr.append(sh)
    mar = OxmlElement("w:tcMar")
    for edge in ("top", "left", "bottom", "right"):
        e = OxmlElement("w:" + edge); e.set(qn("w:w"), "70"); e.set(qn("w:type"), "dxa"); mar.append(e)
    tcpr.append(mar)


def fill_cell(cell, text, header):
    p = cell.paragraphs[0]
    ppr = p._p.get_or_add_pPr()
    sp = OxmlElement("w:spacing"); sp.set(qn("w:after"), "40"); sp.set(qn("w:line"), "240")
    sp.set(qn("w:lineRule"), "auto"); ppr.append(sp)
    r = OxmlElement("w:r"); rpr = copy.deepcopy(RUN_RPR)
    if header:
        rf = rpr.find(qn("w:rFonts"))
        (rf.addnext if rf is not None else (lambda e: rpr.insert(0, e)))(OxmlElement("w:b"))
    r.append(rpr)
    t = OxmlElement("w:t"); t.text = text; t.set(qn("xml:space"), "preserve"); r.append(t)
    p._p.append(r)


def keep_next(p_el):
    ppr = p_el.find(qn("w:pPr"))
    if ppr is None:
        ppr = OxmlElement("w:pPr"); p_el.insert(0, ppr)
    ps = ppr.find(qn("w:pStyle"))
    kn = OxmlElement("w:keepNext")
    (ps.addnext if ps is not None else (lambda e: ppr.insert(0, e)))(kn)
    return p_el


def TAB(headers, rows, widths):
    t = d.add_table(rows=1, cols=len(headers))
    t.autofit = False
    tblpr = t._tbl.tblPr
    lay = OxmlElement("w:tblLayout"); lay.set(qn("w:type"), "fixed")
    look = tblpr.find(qn("w:tblLook"))
    (look.addprevious if look is not None else tblpr.append)(lay)
    for c, w in zip(t.columns, widths):
        c.width = Inches(w)
    hdr = t.rows[0]
    trpr = hdr._tr.get_or_add_trPr()
    trpr.append(OxmlElement("w:tblHeader")); trpr.append(OxmlElement("w:cantSplit"))
    for c, txt, w in zip(hdr.cells, headers, widths):
        cell_fmt(c, True, w); fill_cell(c, txt, True)
    for row in rows:
        r = t.add_row()
        r._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for c, txt, w in zip(r.cells, row, widths):
            cell_fmt(c, False, w); fill_cell(c, str(txt), False)
    # giu hang tieu de voi hang dau; bang ngan thi giu ca bang tren mot trang
    all_rows = list(t.rows)
    keep_rows = all_rows[:-1] if len(rows) <= 5 else all_rows[:1]
    for r in keep_rows:
        for c in r.cells:
            keep_next(c.paragraphs[0]._p)
    prev = t._tbl.getprevious()
    if prev is not None and prev.tag == qn("w:p") and ptext(prev).strip():
        keep_next(prev)
    # doan trong sau bang (giong ban goc)
    gap = OxmlElement("w:p"); ppr = OxmlElement("w:pPr"); sp = OxmlElement("w:spacing")
    sp.set(qn("w:after"), "0"); ppr.append(sp); gap.append(ppr); add(gap)
    return t


def set_cell(tbl_el, r, c, text):
    tbl = docx.table.Table(tbl_el, d)
    set_runs_text(tbl.rows[r].cells[c].paragraphs[0]._p, text)


def clone_row(tbl_el, src_idx, after_idx, texts):
    trs = tbl_el.findall(qn("w:tr"))
    new = copy.deepcopy(trs[src_idx])
    trs[after_idx].addnext(new)
    for tc, text in zip(new.findall(qn("w:tc")), texts):
        set_runs_text(tc.find(qn("w:p")), text)


# ======================= Muc 1-6: sua tai cho =======================
h1 = find_p("Yêu cầu thiết kế")
numpr = h1.find(qn("w:pPr")).find(qn("w:numPr"))
if numpr is not None:
    numpr.getparent().remove(numpr)
set_runs_text(h1, "1 Yêu cầu thiết kế")

sow = tables()[0]
set_cell(sow, 1, 2, "12 222 kg")
cell = docx.table.Table(sow, d).rows[2].cells[1].paragraphs[0]._p
for r in cell.findall(qn("w:r")):
    t = r.find(qn("w:t"))
    if t is not None and t.text == "lm":
        t.text = "h"
set_cell(sow, 4, 2, "0,12 m/s")
set_cell(sow, 7, 2, "0,25 m")
set_cell(sow, 8, 2, "0,2")
set_cell(sow, 9, 2, "0,96")
set_cell(sow, 10, 2, "1,2")
set_cell(sow, 11, 2, "3 pha 380/400 V – 50 Hz")
note = mk(T_P, "Thông số chọn thêm theo bảng mẫu x = 1: tỷ số truyền i = 100; quán tính rotor "
          "J_rot = 0,039 kg·m²; g = 9,81 m/s². Khối lượng chuyển động có tải m = Q + m_h = 12 444 kg, "
          "không tải m₀ = m_h = 222 kg. Chiều cao nâng h = 12 m không phải hành trình ngang. Ô khối lượng "
          "bổ sung E51 của bảng tính để 0 vì chưa có khối lượng khung xe; phải nhập số thực khi chốt thiết kế.")
after_sow = sow.getnext()
(after_sow if after_sow is not None and after_sow.tag == qn("w:p") and not ptext(after_sow).strip()
 else sow).addnext(note)

h2 = find_p("Lực kéo và quy đổi về trục động cơ")
h2_ppr = h2.find(qn("w:pPr"))
h2.replace(h2_ppr, copy.deepcopy(T_H1.find(qn("w:pPr"))))
run = h2.find(qn("w:r"))
run.replace(run.find(qn("w:rPr")), copy.deepcopy(T_H1.find(qn("w:r")).find(qn("w:rPr"))))
for br in h2.iter(qn("w:lastRenderedPageBreak")):
    br.getparent().remove(br)
set_runs_text(h2, "2 Lực kéo và quy đổi về trục động cơ")

red = find_p("Với")
for col in list(red.iter(qn("w:color"))):
    col.getparent().remove(col)

rotor = mk(T_P, "Quán tính rotor quy về chuyển động thẳng: m_rot = J_rot (2i/D)² = 0,039 × 800² = 24 960 kg, "
           "gấp 2,0 lần khối lượng có tải và 112 lần khối lượng không tải. Vì vậy mô-men tăng tốc rotor chiếm "
           "15,6 N·m trong 55,5 N·m ở a = 0,5 (trước dự trữ); động năng rotor 179,7 J lớn hơn động năng xe có tải "
           "89,6 J; lượt không tải có tái sinh năng lượng khi giảm tốc (mục 7.2) và, nếu phanh không tác động, "
           "xe không tải trôi xa hơn nhiều so với xe có tải (mục 7.4).")
find_p("J_tải là quán tính quy đổi động học").addnext(rotor)

set_runs_text(find_p("Hình 1."),
              "Hình 1. Chu kỳ có tải đi và không tải về. F là lực tại bánh; P và M tại trục motor. Chiều về có v, F, "
              "M âm. Cuối lượt không tải, mô-men motor đổi dấu do quán tính rotor nên P âm nhẹ (đỉnh ≈ −21 W, "
              "≈ 50 J): motor hãm tái sinh. Lượt có tải không tái sinh vì lực cản lớn hơn lực quán tính.")
set_runs_text(find_p("Bảng tính tích phân bình phương theo từng pha"),
              "Bảng tính tích phân bình phương theo từng pha, kể cả nghỉ với P = M = 0. Trong mỗi pha jerk không "
              "đổi, P² là đa thức bậc tối đa 6 nên tích phân Gauss 4 điểm cho kết quả chính xác theo mô hình. "
              "Không lấy bình phương công suất trung bình của pha để thay cho trung bình bình phương công suất.")

set_runs_text(find_p("Chọn sơ bộ motor không đồng bộ"),
              "Chọn motor không đồng bộ ba pha 5,5 kW, 6 cực, mã ABB M3AA 132ME 6 (dòng Process performance, "
              "IE3, mã sản phẩm 3GAA133350-…; dòng định mức 12,0 A tại 400 V theo nhà phân phối [9]). Đặt hàng "
              "thêm đầu trục thứ hai (NDE) để lắp phanh và encoder, và cảm biến nhiệt PTC trong cuộn dây. Các trị "
              "số nhãn khác lấy theo sheet x = 1; cần datasheet chính thức của đúng mã đặt hàng để chốt.")
chk = [t for t in tables() if ptext(t).startswith("Kiểm tra")][0]
set_cell(chk, 6, 1, "66,59 N·m = 1,235 T_đm < T_max = 156 N·m; dòng ≈ 13,6 A < I_N = 17 A")
set_cell(chk, 6, 2, "Đạt; không khai thác do giới hạn 47,9 N·m")
clone_row(chk, 4, 4, ["Dòng VFD cho quỹ đạo", "≈ 10,6 A ≤ I_Hd = 12,6 A", "Đạt"])
set_runs_text(find_p("Mô-men khởi động trực tiếp 107,8"),
              "Khi chạy biến tần, mô-men khả dụng do giới hạn dòng/mô-men của VFD quyết định; không dùng mô-men "
              "khởi động trực tiếp 107,8 N·m. Ước tính dòng stator ở từ thông định mức theo "
              "I = √[(I_đm sinφ)² + (I_đm cosφ · M/M_đm)²]: mô-men quỹ đạo có dự trữ 42,05 N·m cần ≈ 10,6 A; bao "
              "66,59 N·m cần ≈ 13,6 A, nhỏ hơn I_N = 17 A của biến tần chọn ở mục 7.1; motor có T_max = 2,9 T_đm "
              "= 156 N·m. Điều kiện quá tải 1,235 lần vì vậy được đáp ứng mà không cần dùng khả năng quá tải của "
              "VFD. Tuy nhiên VFD đặt giới hạn mô-men 47,9 N·m để bảo vệ hộp số (mục 6), nên vận hành không khai "
              "thác a = 0,5.")

mech = [t for t in tables() if ptext(t).startswith("Hạng mục")][-1]
set_cell(mech, 2, 1, "i = 100; η = 0,96; n₁ = 917 vòng/phút, n₂ = 9,17 vòng/phút; P₁ = 5,5 kW; "
                     "M₂ cho phép ≥ 4,6 kN·m (hộp côn–trụ 3 cấp).")
clone_row(mech, 4, 4, ["Giới hạn mô-men VFD", "47,9 N·m tại trục motor → 47,9 × 100 × 0,96 = 4,6 kN·m ở đầu ra."])
set_runs_text(find_p("Motor có thể phát mô-men lớn hơn nhu cầu kéo"),
              "Motor có thể phát tới T_max = 2,9 T_đm = 156 N·m, tương ứng ≈ 15 kN·m ở đầu ra hộp số, lớn hơn "
              "nhiều mức tính toán. Vì vậy đặt giới hạn mô-men trong VFD M_lim = k F D/(2iη) = 47,9 N·m "
              "(≈ 0,89 T_đm) để mô-men đầu ra không vượt 4,6 kN·m. Mức này vẫn lớn hơn 1,37 lần mô-men đỉnh quỹ "
              "đạo 35,04 N·m; với tải đầy, gia tốc lớn nhất đạt được ≈ 0,34 m/s². Khi kẹt bánh, giới hạn mô-men "
              "VFD chặn lực tác dụng lên hộp số và khớp nối.")

# ======================= Xoa muc 7 tro di va dung lai =======================
start = find_p("7 Biến tần, bảo vệ và hãm")
keep_text = {k: ptext(find_p(k)) for k in (
    "Ngoài các biểu thức trên", "Chạm giới hạn vận hành bên phải", "Giới hạn tải của cơ cấu nâng",
    "ZERO_SPEED khởi tạo")}
steps = [ptext(find_p(f"{n}. ")) for n in range(1, 7)]
el = start
while el is not None and el.tag != qn("w:sectPr"):
    nxt = el.getnext(); body.remove(el); el = nxt

# ---------------- 7 ----------------
H1("7 Biến tần, bảo vệ và hãm")
P("Mục này chốt biến tần, điện trở hãm, thiết bị đóng cắt, cáp và phanh theo tải đã tính ở mục 2–5. Số liệu "
  "thiết bị lấy từ catalog hãng và nhà phân phối tra ngày 24/09/2026 (mục 14).")
H2("7.1 Chọn biến tần")
P("Chọn ABB ACS380-042C-17A0-4 (dòng biến tần máy ACS380, kế nhiệm ACS355 trong tài liệu mẫu) với mã tùy "
  "chọn +K492 (PROFINET, FPNO-21) và +L535 (module encoder BTAC-02, HTL/TTL). Biến tần điều khiển vector có "
  "encoder, có STO và chopper hãm tích hợp [5].")
TAB(["Tiêu chí", "Yêu cầu", "ACS380-…-17A0-4", "Đánh giá"], [
    ["Công suất (SOW)", "≥ 1,3 × 5,5 = 7,15 kW", "P_N = 7,5 kW", "Đạt"],
    ["Dòng tải nặng ≥ dòng motor", "≥ 12 A", "I_Hd = 12,6 A (5,5 kW)", "Đạt"],
    ["Dòng định mức", "≥ 12 A", "I_N = 17 A", "Đạt, dư 42 %"],
    ["Dòng ở mô-men quỹ đạo có dự trữ", "≈ 10,6 A", "≤ I_Hd", "Đạt"],
    ["Dòng ở bao a = 0,5", "≈ 13,6 A", "≤ I_N; I_max = 22,7 A", "Đạt, không cần quá tải"],
    ["Phản hồi tốc độ", "Encoder HTL 15,6 kHz", "BTAC-02", "Đạt"],
    ["An toàn", "STO 2 kênh", "STO tích hợp", "Đạt"],
], [2.0, 1.75, 1.75, 1.5])
EQ("I ≈ √[(I_đm sinφ)² + (I_đm cosφ · M/M_đm)²],  I_đm = 12 A, cosφ = 0,74")
P("Ước tính dòng giả thiết từ thông định mức và bỏ qua bão hòa; kiểm tra lại bằng dòng đọc trên biến tần khi "
  "chạy thử. Cài đặt chính: dữ liệu nhãn motor (400 V; 12 A; 973 vòng/phút; 50 Hz; 5,5 kW; cosφ 0,74); điều "
  "khiển vector có encoder 1 024 xung/vòng; giới hạn mô-men 47,9 N·m; giới hạn dòng 18 A (1,5 I_đm, nhỏ hơn "
  "I_max = 22,7 A); tốc độ tối đa 1 008 vòng/phút (110 %); dừng theo ramp khi mất truyền thông. Phanh do PLC "
  "điều khiển theo lưu đồ S0–S8, không bật đồng thời chức năng điều khiển phanh của VFD.")

H2("7.2 Năng lượng hãm và điện trở hãm")
P("Với μ = 0,2, lực cản của xe có tải (24,4 kN) lớn hơn lực quán tính khi giảm tốc kể cả ở a = 0,5 (6,2 kN), "
  "nên lượt có tải không trả năng lượng về DC bus. Lượt không tải có tái sinh vì động năng rotor 179,7 J lớn "
  "hơn nhiều động năng xe không tải 1,6 J. Kết quả tích phân số theo quỹ đạo:")
TAB(["Trường hợp", "Công suất tái sinh đỉnh", "Năng lượng mỗi lần dừng"], [
    ["Có tải, mọi quỹ đạo a ≤ 0,5", "0", "0"],
    ["Không tải, giảm tốc 5 s theo chu kỳ", "21 W", "50 J"],
    ["Không tải, dừng nhanh nhất cho phép (j = 2; a = 0,49 m/s²; 0,49 s)", "0,77 kW", "168 J"],
    ["Bao trên: a = 0,5 ngay tại v_max, không tải", "1,46 kW", "≤ 180 J (động năng rotor)"],
], [3.6, 1.7, 1.7])
EQ("P_hãm = [J_rot (2i/D) a − m₀ (μg − a) D/(2iη)] ω_motor")
P("ACS380-17A0-4 cho phép điện trở hãm R_min = 32 Ω và R_max = 54 Ω; khả năng hãm P_BRcont = 5,5 kW, "
  "P_BRmax = 8,25 kW [5]. Chọn R_h = 39 Ω, theo điện trở ví dụ ABB nêu cho cỡ này (Danotherm CBR-V 560 D HT "
  "406 39R). Công suất hãm lớn nhất 1,46 kW nhỏ hơn nhiều 8,25 kW; năng lượng mỗi lần dừng ≤ 0,18 kJ; công suất "
  "trung bình theo chu kỳ dưới 1 W. Cần kiểm tra năng lượng xung cho phép của điện trở ≥ 0,2 kJ trên datasheet. "
  "Không dùng 200 Ω của tài liệu mẫu vì nằm ngoài dải 32–54 Ω của biến tần đã chọn.")

H2("7.3 Thiết bị đóng cắt, bảo vệ đầu vào và cáp")
TAB(["Hạng mục", "Tiêu chí", "Lựa chọn", "Kiểm tra"], [
    ["MCCB Q1", "I_bv ≥ 1,5 × 12 = 18 A (SOW); không lớn hơn cầu chì gG 32 A ABB khuyến nghị",
     "Schneider EZC100N3020: 3P 20 A; Icu 15 kA/415 V", "Đạt; dòng vào ước tính ≈ 15 A ở tải định mức motor"],
    ["Contactor lực K1", "Cắt nguồn lực khi dừng khẩn (SOW); AC-1 ≥ dòng vào",
     "Schneider LC1D18BD: 18 A AC-3; 32 A AC-1; cuộn 24 VDC", "Đạt; K1 không đóng cắt theo mỗi lần chạy"],
    ["Cáp nguồn vào VFD", "I_z ≥ I_n Q1 = 20 A", "Cu/PVC 4×4 mm²: I_z = 32 × 0,87 = 27,8 A", "Đạt"],
    ["Cáp motor", "I_z ≥ 12 A; cáp đối xứng có màn chắn",
     "Cu 3×2,5 mm² + PE có màn chắn: I_z = 24 × 0,87 = 20,9 A", "Đạt; ΔU ≈ 0,9 % với 25 m"],
    ["Cáp phanh / encoder", "24 VDC; tín hiệu HTL",
     "2×1,5 mm²; 4×2×0,25 mm² xoắn đôi có màn chắn", "Đi máng riêng với cáp motor"],
], [1.3, 2.0, 2.1, 1.6])
P("Dòng cho phép lấy theo IEC 60364-5-52, bảng B.52.4 (PVC, 3 dây mang tải, lắp đặt kiểu C), hệ số nhiệt độ "
  "40 °C là 0,87. Dòng vào VFD ước tính với hệ số công suất đầu vào ≈ 0,6. Chiều dài 25 m là giả thiết cho cáp "
  "lên xe con; thay bằng chiều dài thực. Nếu dòng ngắn mạch tại tủ lớn hơn 10 kA, ABB khuyến nghị lắp cuộn kháng "
  "đầu vào cho cỡ R3 [5]; Icu của Q1 phải lớn hơn dòng ngắn mạch thực tại tủ.")

H2("7.4 Phanh cơ")
P("Trong vận hành thường, phanh chỉ đóng ở 0 tốc độ (S6); hãm động chỉ xảy ra khi dừng khẩn hoặc mất điều "
  "khiển. Vì μ = 0,2 làm xe có tải tự dừng rất nhanh, cỡ phanh do trường hợp không tải quyết định, nơi quán tính "
  "rotor chiếm gần hết.")
EQ("a_phanh = −[M_b + m μ g D/(2iη)] / [J_rot (2i/D) + m D/(2iη)]")
TAB(["Đại lượng (M_b = 32 N·m, trễ tác động 0,1 s)", "Có tải", "Không tải"], [
    ["Gia tốc khi phanh", "1,35 m/s²", "1,03 m/s²"],
    ["Thời gian / quãng đường dừng", "0,09 s / 17 mm", "0,12 s / 19 mm"],
    ["STO tác động nhưng phanh không đóng", "0,67 m/s²; 23 mm", "0,018 m/s²; 0,41 m trong 6,7 s"],
    ["Mô-men đầu ra hộp số khi phanh", "0,96 kN·m", "≈ 0,03 kN·m"],
], [3.4, 1.8, 1.8])
P("Chọn phanh lò xo INTORQ BFK458-12E, 24 VDC: mô-men phanh làm việc 32 N·m, tốc độ tối đa 3 600 vòng/phút, "
  "công ma sát cho phép một lần 24 kJ [8]. Cỡ 12 là cỡ nhỏ nhất có microswitch giám sát trạng thái đóng/mở, cần "
  "cho tín hiệu BRAKE_OPEN/BRAKE_CLOSED. Năng lượng mỗi lần dừng khẩn ≤ 269 J, nhỏ hơn nhiều 24 kJ. Kiểu E cho "
  "phép giảm mô-men bằng vòng chỉnh nếu thử nghiệm cho thấy tải lắc quá mức. Gia tốc dừng khẩn 1,0–1,35 m/s² "
  "vượt 0,5 m/s²; điều này chấp nhận cho dừng khẩn vì giới hạn a, j của đề áp dụng cho chuyển động có điều "
  "khiển.")

H2("7.5 Khoảng cách dừng và vị trí công tắc")
EQ("s_dừng,có tải = 0,5 × 0,12 × 10 = 0,60 m")
EQ("d_kích hoạt ≥ v_max t_trễ + s_dừng + Δs")
P("Với tổng trễ 0,10 s và dự trữ 0,05 m: d ≥ 0,662 m; đặt công tắc hành trình vận hành SQ1/SQ2 cách điểm dừng "
  "0,70 m. Không tải giảm tốc 5 s chỉ cần 0,30 m. Stop giữa ramp phải tính quỹ đạo từ trạng thái hiện tại.")
P("Cực hạn SQ3/SQ4 đặt sau điểm dừng vận hành 0,10 m. Khoảng từ cực hạn đến chặn cơ khí/giảm chấn phải lớn hơn "
  "quãng đường dừng khẩn có phanh (≈ 0,02 m) và nên bao cả trường hợp phanh không tác động khi không tải "
  "(0,41 m) cộng dự trữ 0,05 m: chọn ≥ 0,50 m. Như vậy vùng vượt hành trình mỗi đầu khoảng 0,6 m sau điểm dừng; "
  "các khoảng cách này phải được đo lại khi chạy thử.")

# ---------------- 8 ----------------
H1("8 Cấu trúc phần cứng, điều khiển FOC và cảm biến")
P("Mạch động lực: lưới 3 pha → Q1 (MCCB) → K1 (contactor lực) → U1 (biến tần) → M1 (motor, phanh, encoder) → "
  "hộp số → bánh xe. PLC tạo quỹ đạo, xử lý lưu đồ và liên khóa; rơ-le an toàn KS xử lý E-stop và cực hạn, tác "
  "động STO hai kênh, cắt điện phanh và cuộn K1 độc lập với PLC.")
PIC("rev_hardware.png", 6.9)
CAP("Hình 2. Sơ đồ khối phần cứng: mạch động lực, điều khiển và nhánh an toàn độc lập. Sơ đồ khối không thay "
    "cho bản vẽ đấu dây.")
H2("Cấu trúc điều khiển FOC hai mạch vòng PI")
P("Biến tần thực hiện FOC định hướng từ thông rotor gián tiếp. Mạch vòng dòng (bên trong) có hai bộ PI điều "
  "chỉnh i_sd (từ thông) và i_sq (mô-men). Mạch vòng tốc độ (bên ngoài) dùng PI tạo i_sq* từ sai lệch ω* − ω, có "
  "giới hạn mô-men 47,9 N·m. PLC chỉ tạo quỹ đạo S-curve v*(t) và gửi tốc độ đặt qua PROFINET; encoder cấp ω "
  "cho vòng tốc độ và mô hình từ thông.")
PIC("rev_foc.png", 6.9)
CAP("Hình 3. Cấu trúc FOC với hai mạch vòng PI nối tầng.")
P("Hằng số thời gian rotor ước tính trong sheet 20251_IM_Params_Calculation của bảng tính khóa trước (cùng loại "
  "motor 5,5 kW, 6 cực) là T_r ≈ 0,20 s. Thời gian từ hóa trước khi mở phanh lấy ≈ 5T_r = 1,0 s, đây là timeout "
  "của trạng thái S2. Chỉnh định tham số PI và mô phỏng MATLAB/Simulink thuộc giai đoạn 4 của SOW.")
H2("Cảm biến phản hồi")
TAB(["Cảm biến", "Lựa chọn", "Thông số", "Kiểm tra"], [
    ["Tốc độ", "Autonics E50S8-1024-3-T-24 → BTAC-02", "1 024 xung/vòng; A/B/Z; totem-pole 12–24 VDC; 300 kHz",
     "917 vòng/phút → 15,6 kHz; 1,9 µm/xung (×4); ngưỡng 0,002 m/s ↔ 261 Hz; IP50 nên lắp dưới nắp che"],
    ["Dòng stator", "Cảm biến dòng trong ACS380; đo kiểm ngoài: LEM LA 55-P",
     "LA 55-P: 50 A; DC–200 kHz; Hall vòng kín", "Đủ cho 12 A định mức, 22,7 A cực đại; biến dòng 50 Hz không đo "
     "được dòng tần số thay đổi tới 0 Hz"],
    ["Hành trình vận hành L/R (SQ1/SQ2)", "Omron D4N-4120", "Cần lăn; 1NC/1NO; cơ cấu mở cưỡng bức",
     "Dùng tiếp điểm NC; đặt 0,70 m trước điểm dừng"],
    ["Cực hạn L/R (SQ3/SQ4)", "Omron D4N-4120, cơ cấu tác động riêng", "Như trên",
     "NC nối vào KS; đặt 0,10 m sau điểm dừng"],
    ["Nhiệt motor", "PTC trong cuộn dây (tùy chọn khi đặt motor)", "Theo cấp cách điện F", "Nối đầu vào nhiệt VFD → TEMP_OK"],
    ["Trạng thái phanh", "Microswitch của BFK458-12", "Tiếp điểm giám sát đóng/mở",
     "BRAKE_OPEN, BRAKE_CLOSED; cấm hai tín hiệu cùng mức"],
], [1.25, 1.75, 1.9, 2.1])

# ---------------- 9 ----------------
H1("9 Tín hiệu và các điều kiện liên khóa")
TAB(["Nhóm", "Tín hiệu tối thiểu", "Ý nghĩa"], [
    ["Lệnh người vận hành", "FWD, REV, STOP, RESET", "Lệnh tiến/lùi kiểu giữ để chạy; Reset cạnh lên riêng."],
    ["An toàn độc lập", "ESTOP_CH1/2, FINAL_L/R, K1_FB, SAFETY_OK",
     "E-stop, cực hạn hai đầu vào KS; phản hồi K1; đầu ra bán dẫn của KS về PLC."],
    ["Hành trình vận hành", "LIMIT_L_OK, LIMIT_R_OK", "Cho phép chạy theo hướng; tiếp điểm NC, kiểm tra đứt dây."],
    ["Biến tần (PROFINET)", "DRIVE_READY, FAULT, TORQUE_READY", "Điều kiện phát mô-men; lỗi mất điều khiển phản ứng ngay."],
    ["Cơ khí và tải", "BRAKE_OPEN, BRAKE_CLOSED, LOAD_OK", "Trạng thái phanh có kiểm tra hợp lý; tải trong ngưỡng hệ cân."],
    ["Giám sát", "SPEED_VALID, ZERO_SPEED, TEMP_OK, OVERSPEED, SPEED_DEV",
     "Encoder hợp lệ; đứng yên; nhiệt; quá tốc 110/120 %; sai lệch tốc độ."],
    ["Đầu ra điều khiển", "ENABLE, SPEED_REF, BRAKE_RELEASE, SAFE_STOP_REQ",
     "Lệnh drive và tốc độ qua PROFINET; nhả phanh; PLC yêu cầu KS dừng bảo vệ."],
    ["Đầu ra thông báo", "FAULT_LATCH, BUZZER, STATUS", "Mã lỗi lưu giữ và cảnh báo trạng thái."],
], [1.5, 2.5, 3.0])
EQ("CMD_VALID = FWD XOR REV")
EQ("DIR_OK = (FWD AND LIMIT_R_OK) OR (REV AND LIMIT_L_OK)")
EQ("PERMIT = SAFETY_OK AND K1_FB AND DRIVE_READY AND TEMP_OK")
EQ("             AND LOAD_OK AND SPEED_VALID AND BRAKE_CLOSED AND NOT FAULT_LATCH")
for k in ("Ngoài các biểu thức trên", "Chạm giới hạn vận hành bên phải", "Giới hạn tải của cơ cấu nâng",
          "ZERO_SPEED khởi tạo"):
    P(keep_text[k])
P("Số đầu vào số khoảng 11 (FWD, REV, STOP, RESET, LIMIT_L/R, BRAKE_OPEN/CLOSED, K1_FB, SAFETY_OK, LOAD_OK), "
  "nằm trong 14 DI của CPU 1214C; đầu ra khoảng 6 (BRAKE_RELEASE, SAFE_STOP_REQ, 3 đèn, còi) trong 10 DO. Lệnh "
  "và trạng thái biến tần trao đổi qua PROFINET. SAFE_STOP_REQ là tiếp điểm PLC nối tiếp trong mạch vào của KS để "
  "PLC kích hoạt cùng nhánh dừng; tiếp điểm này không thay cho chức năng an toàn.")

# ---------------- 10 ----------------
H1("10 Lưu đồ trạng thái vận hành")
PIC("rev_states.png", 6.9)
CAP("Hình 4. Lưu đồ trạng thái S0–S8. Nhánh S8 có hiệu lực từ mọi trạng thái S0–S7 và do rơ-le an toàn thực "
    "hiện độc lập với PLC. S7 tự về S0 khi nguyên nhân đã hết; Reset chỉ nhấn một lần tại S0.")

# ---------------- 11 ----------------
H1("11 Bảng chuyển trạng thái và xử lý lỗi")
TAB(["Trạng thái", "Tác động", "Điều kiện chuyển"], [
    ["S0 Khóa khởi động", "STO tác động; phanh đóng; K1 đóng khi KS đã reset",
     "Tự kiểm tra OK + SAFETY_OK + không lỗi + RESET↑ + mọi lệnh chạy đã nhả → S1"],
    ["S1 Sẵn sàng", "Phanh đóng; mô-men 0",
     "Lệnh mới FWD⊕REV (cạnh lên) + PERMIT + DIR_OK → S2; lệnh giữ từ trước không có hiệu lực"],
    ["S2 Tạo mô-men", "Bỏ STO; ENABLE; ω* = 0; từ hóa",
     "TORQUE_READY → S3; nhả lệnh → S1 (STO lại); quá 1,0 s → S7"],
    ["S3 Mở phanh", "BRAKE_RELEASE = 1; giữ 0 tốc độ",
     "BRAKE_OPEN → S4; nhả lệnh, mất PERMIT hoặc quá 0,5 s → S6 (quá thời gian được ghi thành lỗi chờ)"],
    ["S4 Chạy", "Bám S-curve; giám sát tốc độ, sai lệch, giới hạn",
     "STOP, nhả nút, đổi chiều, giới hạn vận hành, lỗi còn điều khiển → S5"],
    ["S5 Giảm tốc", "S-curve về 0 (có tải 10 s, không tải 5 s)",
     "ZERO_SPEED (|v| < 0,002 m/s trong 0,2 s) → S6; quá 12 s → S8"],
    ["S6 Đóng phanh", "BRAKE_RELEASE = 0; giữ 0 tốc độ",
     "BRAKE_CLOSED trong 0,5 s → STO → S1, hoặc S7 nếu có lỗi chờ; quá 0,5 s → STO → S7 (lỗi phanh)"],
    ["S7 Lỗi đã chốt", "STO; phanh đóng; cấm chạy; lưu mã lỗi", "Nguyên nhân hết + SAFETY_OK → S0 (tự động)"],
    ["S8 Dừng khẩn / bảo vệ", "KS: STO 2 kênh + cắt điện phanh + mở K1",
     "Từ mọi trạng thái khi E-stop, cực hạn, lỗi mạch an toàn, VFD trip, mất encoder, mất nguồn, quá tốc > 120 %, "
     "timeout S5. Đứng yên + chốt lỗi → S7; nhả E-stop không tự chạy lại"],
], [1.35, 2.2, 3.45])
P("Khác cơ cấu nâng, xe con chạy ngang không có tải trọng trường kéo trôi. Vì vậy khi BRAKE_CLOSED không xác "
  "nhận trong 0,5 s, hệ bỏ mô-men rồi chốt lỗi phanh; cấm chạy cho tới khi kiểm tra phanh. Timeout S5 lấy 12 s, "
  "lớn hơn thời gian giảm tốc dài nhất 10 s.")
H2("Bảo vệ quá áp, quá dòng, quá nhiệt và giám sát")
TAB(["Bảo vệ", "Phát hiện", "Ngưỡng cài đặt thử", "Phản ứng"], [
    ["Quá dòng tức thời", "VFD (phần cứng)", "Theo VFD", "Khóa IGBT, trip → S8 → S7"],
    ["Quá tải dòng, mô-men", "VFD", "Giới hạn dòng 18 A; mô-men 47,9 N·m; I²t motor với I_đm = 12 A",
     "Giới hạn liên tục; cảnh báo → S5 → S7; trip → S8"],
    ["Quá áp DC bus", "VFD", "Chopper xả vào R_h = 39 Ω; trip theo ngưỡng VFD", "Trip → S8 → S7"],
    ["Thấp áp, mất pha, mất nguồn", "VFD, K1_FB", "Theo VFD", "S8 → S7"],
    ["Quá nhiệt motor", "PTC + mô hình nhiệt VFD", "Cảnh báo trước; trip khi PTC tác động",
     "Cảnh báo: hoàn tất dừng S5 → S6 → S7; trip → S8"],
    ["Quá nhiệt biến tần", "Cảm biến tản nhiệt VFD", "Theo VFD", "Cảnh báo → S5; trip → S8"],
    ["Quá tốc", "Encoder; PLC và VFD", "110 % (0,132 m/s) trong 0,1 s; 120 % (0,144 m/s)", "110 %: S5 → S7; 120 %: S8"],
    ["Sai lệch tốc độ", "PLC", "|v* − v| > 0,024 m/s (20 % v_max) trong 0,5 s", "S5 → S7; không giảm được → S8"],
    ["Mất encoder", "BTAC-02, VFD", "Đứt dây, mất xung", "S8 → S7"],
    ["Mất truyền thông", "VFD và PLC", "Watchdog PROFINET", "VFD dừng theo ramp → S5 → S7"],
    ["Phanh không mở / không đóng", "Microswitch phanh", "0,5 s", "S3 → S6 → S7 / S6 → S7"],
    ["Giới hạn vận hành", "SQ1/SQ2 (NC)", "Tiếp điểm mở", "S4 → S5; cấm tiếp chiều đó"],
    ["Cực hạn, E-stop", "SQ3/SQ4, E-stop → KS", "Tiếp điểm mở", "S8"],
    ["Quá tải nâng", "LOAD_OK", "Theo hệ cân tải", "S5 → S7; khóa chạy mới"],
], [1.55, 1.45, 2.25, 1.75])
P("Lỗi còn điều khiển được đi S5 → S6 → S7: giảm tốc theo S-curve, đóng phanh rồi mới bỏ mô-men. Lỗi mất điều "
  "khiển đi thẳng S8, không chờ S-curve hoặc TORQUE_READY.")
P("Dừng khẩn thực hiện theo loại 0 của IEC 60204-1: KS cắt đồng thời STO hai kênh, nguồn nhả phanh và cuộn K1. "
  "STO ngăn tạo mô-men ngay lập tức; K1 cắt nguồn lực của biến tần, đáp ứng yêu cầu SOW “E-stop ngắt nguồn lực "
  "ngay lập tức”. STO không thay cho cách ly khi bảo trì; khi bảo trì dùng Q1 có khóa. Nếu đánh giá rủi ro yêu "
  "cầu dừng loại 1, phải dùng SS1 với timer an toàn được thẩm định rồi mới STO [6–7].")

# ---------------- 12 ----------------
H1("12 Quy trình vận hành và kiểm chứng")
H2("Trình tự vận hành")
steps[1] = ("2. Cấp nguồn: đóng Q1; reset KS để đóng K1; mạch ở S0, phanh đóng. Kiểm tra E-stop và cực hạn hợp "
            "lệ, lệnh tiến/lùi đã nhả; nhấn Reset một lần để sang S1.")
steps[4] = ("5. Sự cố: dùng E-stop khi cần dừng khẩn. Sau dừng, xử lý nguyên nhân và kiểm tra phanh; nhả E-stop, "
            "reset KS; khi hết lỗi hệ tự về S0, nhấn Reset rồi ra lệnh chạy mới. Không nối tắt liên khóa.")
for s in steps:
    P(s)
H2("Ma trận kiểm chứng trước nghiệm thu")
TAB(["Phép kiểm tra", "Kết quả cần đạt"], [
    ["Chạy có tải / không tải", "Đo |v| ≤ 0,12 m/s; |a| ≤ 0,5 m/s²; |j| ≤ 2 m/s³ trong chuyển động thường; nêu cách lọc/ước lượng đạo hàm."],
    ["Dòng và mô-men", "Dòng khi tăng tốc có tải ≈ 10–11 A; không chạm giới hạn 47,9 N·m trong chu kỳ thường."],
    ["Bám ray và phân bố tải", "Không trượt trong điều kiện ray thực; phản lực bánh chủ động đủ ở tải lệch bất lợi."],
    ["Stop khi đang tăng tốc", "Gia tốc không nhảy bậc; quỹ đạo dừng giữ jerk trong giới hạn."],
    ["Hai nút chiều cùng tác động", "Không khởi động; nếu đang chạy thì dừng có điều khiển."],
    ["Hành trình vận hành / cực hạn", "Chỉ rút khỏi giới hạn vận hành theo hướng cho phép; cực hạn khóa cả hai chiều."],
    ["Timeout từ hóa, phanh", "Ngắt microswitch hoặc làm chậm phanh: chốt lỗi đúng 1,0 s (S2) và 0,5 s (S3, S6)."],
    ["Quá tốc, sai lệch tốc độ", "Tín hiệu mô phỏng vượt ngưỡng 110/120 % và 0,024 m/s: vào S5/S8 đúng bảng mục 11."],
    ["E-stop, mất nguồn, VFD trip", "KS tác động độc lập PLC; K1 mở; đo quãng đường dừng có tải và không tải; "
     "khoảng cực hạn → chặn cơ khí ≥ 0,5 m."],
    ["Khôi phục điện và nhả E-stop", "Không tự khởi động lại; cần reset KS, Reset tại S0 và lệnh chạy mới."],
    ["Nhiệt, hãm và chu kỳ lặp", "Dòng, nhiệt motor/VFD/phanh và năng lượng điện trở hãm trong định mức thực."],
], [2.3, 4.7])

# ---------------- 13 ----------------
H1("13 Danh mục thiết bị và vật tư")
TAB(["TT", "Ký hiệu", "Thiết bị", "Mã hiệu – hãng", "Thông số chính", "SL"], [
    ["1", "M1", "Động cơ KĐB 3 pha", "M3AA 132ME 6 (IE3) – ABB",
     "5,5 kW; 400 V Δ; 50 Hz; 6 cực; 12 A; 53,9 N·m; J 0,039 kg·m²; thêm đầu trục NDE, PTC", "1"],
    ["2", "YB", "Phanh lò xo điện từ", "BFK458-12E, 24 VDC, có microswitch – INTORQ (Kendrion)",
     "32 N·m; n_max 3 600 vòng/phút; W_max 24 kJ", "1"],
    ["3", "BG1", "Encoder tăng dần", "E50S8-1024-3-T-24 – Autonics", "1 024 xung; A/B/Z; HTL 12–24 VDC; 300 kHz", "1"],
    ["4", "–", "Hộp giảm tốc", "Côn–trụ 3 cấp (SEW, NORD, Bonfiglioli…; chọn theo catalog)",
     "i ≈ 100; P₁ 5,5 kW; n₁ 917 vòng/phút; M₂ ≥ 4,6 kN·m; η ≥ 0,96", "1"],
    ["5", "U1", "Biến tần", "ACS380-042C-17A0-4 +K492 +L535 – ABB",
     "7,5 kW; I_N 17 A; I_Hd 12,6 A; I_max 22,7 A; STO; chopper; PROFINET; BTAC-02", "1"],
    ["6", "R1", "Điện trở hãm", "39 Ω; ví dụ Danotherm CBR-V 560 D HT 406 39R", "32 ≤ R ≤ 54 Ω; E_xung ≥ 0,2 kJ", "1"],
    ["7", "Q1", "MCCB", "EZC100N3020 – Schneider", "3P 20 A; Icu 15 kA/415 V", "1"],
    ["8", "K1", "Contactor lực", "LC1D18BD – Schneider", "18 A AC-3; 32 A AC-1; cuộn 24 VDC", "1"],
    ["9", "KS", "Rơ-le an toàn", "PNOZ s4 24VDC 3 n/o 1 n/c (750104) – Pilz", "PL e / SIL 3; 2 kênh; 3NO + 1NC + 1 bán dẫn", "1"],
    ["10", "S0", "Nút dừng khẩn", "XB5AS8444 – Schneider", "Ø40, nấm đỏ, xoay nhả; 2NC; IP66", "2"],
    ["11", "SQ1–SQ4", "Công tắc hành trình", "D4N-4120 – Omron", "Cần lăn; 1NC/1NO; mở cưỡng bức", "4"],
    ["12", "A1", "PLC", "CPU 1214C DC/DC/DC (6ES7214-1AG40-0XB0) – Siemens", "14 DI/10 DO/2 AI; PROFINET", "1"],
    ["13", "A2", "HMI", "KTP700 Basic PN (6AV2123-2GB03-0AX0) – Siemens", "Màn hình 7 inch; PROFINET", "1"],
    ["14", "KB", "Contactor phụ cắt phanh", "CAD32BD – Schneider (ví dụ)", "Cuộn 24 VDC; tiếp điểm gương; kiểm tra định mức DC-13 theo cuộn phanh", "2"],
    ["15", "BA", "Cảm biến dòng đo kiểm", "LA 55-P – LEM (tùy chọn)", "50 A; DC–200 kHz", "2"],
    ["16", "G1", "Nguồn 24 VDC", "NDR-120-24 – Mean Well", "24 V, 5 A, gắn DIN", "1"],
    ["17", "W1–W4", "Cáp", "Cu/PVC", "4×4 mm²; 3×2,5 mm² + PE có màn chắn; 2×1,5 mm²; 4×2×0,25 mm² xoắn đôi có màn chắn", "Theo tuyến"],
], [0.35, 0.7, 1.15, 1.95, 2.3, 0.55])
P("Mã hiệu được kiểm tra theo catalog hoặc nhà phân phối ngày 24/09/2026. Hậu tố đặt hàng của motor (kiểu lắp, "
  "điện áp), phụ kiện phanh, điện trở hãm và contactor phụ cần xác nhận lại với nhà cung cấp trước khi mua; có thể "
  "thay bằng thiết bị tương đương nếu đáp ứng cùng thông số.")

# ---------------- 14 ----------------
H1("14 Mức hoàn thiện và tài liệu tham khảo")
P("Báo cáo đã có: phép tính lực, quy đổi mô-men/quán tính, chu kỳ khép kín, đồ thị v–F–P–M, công suất hiệu dụng; "
  "kiểm nghiệm motor; chọn biến tần theo công suất và dòng; điện trở hãm; MCCB, contactor, cáp; phanh và khoảng "
  "cách dừng; sơ đồ khối phần cứng và FOC; cảm biến; lưu đồ trạng thái với timeout và bảng bảo vệ; danh mục thiết bị.")
P("Còn cần trước khi chế tạo: xác nhận khối lượng khung xe và mô hình cản/bám ray (μ = 0,2) với giảng viên; "
  "datasheet chính thức cho mã đặt hàng; chỉnh định PI và mô phỏng MATLAB/Simulink (giai đoạn 4 SOW); bản vẽ đấu "
  "dây và kiểm chứng trên thiết bị. Báo cáo không tuyên bố cơ cấu đạt chứng nhận an toàn.")
H2("Tài liệu dự án")
for t in ["[1] SYSTEM ARCHITECTURE – CRANE TROLLEY DRIVE SYSTEM.txt: tài liệu mẫu tham khảo cấu trúc chức năng; "
          "số liệu trong tệp mẫu không dùng cho thiết kế x = 2.",
          "[2] Tính toán thông số động cơ.xlsx: sheet a=1 là cơ sở cấu hình; sheet 20251_IM_Params_Calculation cho T_r.",
          "[3] SOW_HungNB.docx: phạm vi, bảng đầu vào và tiêu chí nghiệm thu."]:
    P(t)
H2("Tài liệu kỹ thuật tham khảo (tra cứu ngày 24/09/2026)")
refs = [
    "[4] SEW-EURODRIVE, Project Planning for Drives: mô hình lực cản và bám bánh. "
    "https://download.sew-eurodrive.com/download/pdf/10522913.pdf",
    "[5] ABB, Machinery drives ACS380 catalog: định mức 400 V, chopper và điện trở hãm (R_min, R_max), cầu chì, "
    "BTAC-02, STO. https://www.auser.fi/wp-content/uploads/ACS380-katalogi.pdf",
    "[6] IEC 60204-32:2023, thiết bị điện của máy nâng. https://webstore.iec.ch/en/publication/63094",
    "[7] ISO 13850:2015, nguyên tắc chức năng dừng khẩn. https://committee.iso.org/standard/59970.html?browse=ics",
    "[8] Kendrion, INTORQ BFK458 spring-applied brake brochure (10/2022). "
    "https://www.kendrion.com/fileadmin/user_upload/Downloads/Brochures_and_Flyers/Industrial_Brakes_INTORQ/brochure-spring-applied-brake-bfk458-2022.10.pdf",
    "[9] ABB M3AA 132ME 6 (3GAA133350), thông tin nhà phân phối. "
    "https://www.mrosupply.com/electric-motors/6092107_emm13556-s-ap_baldor/",
    "[10] Autonics E50S8-1024-3-T-24. https://www.autonics.com/us/model/E50S8-1024-3-T-24",
    "[11] Pilz PNOZ s4 (750104). https://www.pilz.com/en-INT/eshop/product/750104",
    "[12] Schneider Electric LC1D18BD. https://www.se.com/ca/en/product/LC1D18BD/",
    "[13] Omron D4N-4120. https://industrial.omron.eu/en/products/D4N-4120",
    "[14] IEC 60364-5-52: lựa chọn và lắp đặt hệ thống đi dây, dòng cho phép của cáp.",
]
for t in refs:
    P(t)

d.core_properties.title = "Thiết kế hệ truyền động xe con cầu trục — x = 2"
d.save(str(DST))
print("saved", DST)
